from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re, shutil, tempfile

ROOT=Path.cwd(); FIG=ROOT/'figures'; OUT=ROOT/'四问结果分析与论文插图.docx'
# (section, filename, step heading, explanation)
items=[]
def add(q, fname, heading, text): items.append((q,fname,heading,text))
Q1='问题一 质量评价与领域配比'
for row in [
('fig_01_indicator_distributions','指标方向统一与分布检查','图中汇总22项指标经方向统一和预处理后的分布。偏态及尺度差异说明需要逐指标标准化处理；该图用于检查输入形态，不单独说明质量优劣。'),
('fig_02_corr_cluster_heatmap','22项指标相关结构','相关热图展示指标间冗余与反向关系，聚类只改变显示顺序。主模型依据平方相关计算冗余度并赋权，因此此图解释权重校正的依据，而非把指标另分成中间维度。'),
('fig_03_weights_redundancy22','平方相关冗余赋权','权重图显示22项指标经过平方相关冗余校正后的相对贡献。各指标仍直接进入质量评价，较高权重表示冗余校正后的相对信息量较大，不代表因果重要性。'),
('fig_04_domain_Q_compare','领域质量分布与均值','该图比较抽样领域和扩展数据中的质量分布及均值区间。结果体现领域间存在异质性；置信区间描述抽样不确定性，不能外推成领域质量的因果差异。'),
('fig_05_pearson_conflict_definition','Pearson冲突候选对定义','图中按 r≤−0.30 标记反向相关候选对，并展示最强反向相关关系。阈值是冲突诊断口径；它用于定位指标张力，不改变质量分的构造。'),
('fig_06_conflict_causes','冲突信号的领域与长度分布','图中并列展示领域冲突信号比例、代表性反向指标对及冲突强度与文本长度的关系。它用于描述冲突出现在哪里；领域校准和相关诊断不能证明误判已消除。'),
('fig_07_resolution','指数凹聚合的冲突消解效果','图中比较线性综合分与指数凹聚合分、分数调整分布及β敏感性。主设定β=0.25；结果说明聚合规则会改变冲突样本的综合分，Q不应解释为0—1概率。'),
('fig_08_content_check','未参与评分的内容侧检验','图中比较外部内容指标随质量分位的变化，这些变量没有参与评分构造。其用途是提供关联性检验；观察到的趋势不构成质量评分有效性的因果证明。'),
('fig_09_extended_conflict_check','扩展数据的冲突复核','图中复核扩展来源的反向相关指标对及其分布。扩展结果用于检查冲突结构是否只由A1抽样产生，不把同一轨迹检查点当作独立实验。'),
('fig_10_regmix_overview','配比实验数据总览','图中展示512组训练配方、13个验证域损失和单域份额与损失的关系。份额与损失呈非线性，支持在配比约束下估计混料响应；这仍是观测关系。'),
('fig_11_test_pred_vs_actual','四类模型独立检验','图中比较线性混料、对数线性混料、加性样条和GBM的预测与实测损失。对数线性模型作为关系解释及约束配比求解主模型，GBM作为预测对照；单项预测排名不等于模型选型结论。'),
('fig_12_marginal_effects_forest','对数线性模型边际效应','左侧给出配方域系数及区间，右侧展示边际效应如何随配比位置变化。系数和边际效应描述模型内的损失响应，受配方单纯形约束，不能直接解读为领域因果效应。'),
('fig_13_domain_transfer_effects','固定配比下的局部领域替代','图中计算将少量配比从一域转到另一域的局部预测变化。转移方向必须满足配比总和约束，展示的是给定基准附近的局部替代效应。'),
('fig_14_transfer_matrix','配方域到验证域的迁移矩阵','矩阵展示各配方域变化对不同验证域损失的预测响应，同名域以边框标出。跨域响应不均一，说明总体平均可能掩盖验证域差异；它属于模型迁移诊断。'),
('fig_15_extrapolation_robustness','跨尺度外推稳健性','图中检查代理模型排序、逐验证域排序和域效应排序在不同尺度上的保持情况。尺度外推表现有边界，配方模型适合已覆盖范围内的条件分析，不能据此承诺任意规模下的绝对误差。'),
('fig_16_Q_and_optimal_mixture','质量分与候选最优配比的关系','图中比较领域质量、质量标量解释力及候选配比变化。质量评价与配比拟合是两条独立计算链，当前配比主模型不使用Q；候选配方不是实测最优。'),
('fig_17_length_correction','独特词长度校正','图中对照未校正分位、20组长度条件分位和统一20/80规则。校正旨在减少文本长度对独特词指标的机械影响，不意味着长度偏差已完全消除。'),
('fig_18_length_sensitivity','长度校正敏感性与机械重复压力测试','图中改变分组数并模拟指标机械重复，观察冲突信号变化。该压力测试评估方法稳定性，不含人工有效性标签，不能作为人工准确率证据。'),
('fig_19_format_calibration','格式指标领域校准','图中比较专业性参照与格式指标校准前后的关系，并展示各领域校准效果。校准保持专业性原含义；它是领域可比性处理，不证明潜在误判已消失。'),
('fig_20_conflict_audit','冲突审计汇总','该图汇总冲突候选、样本触发及相关审计结果。诊断阈值仅决定标记范围，不单独改变Q；冲突率不是模型优化目标。')]: add(Q1,*row)

Q2='问题二 共同修正标度律'
for row in [
('fig_21_q2_scale_and_sources','经典规模律与来源校准','该图将经典规模拟合和来源内校准留出表现并列展示。来源基线允许不同，能减少验证语料与训练设置差异造成的混淆；来源校准不等于完整跨来源外推验证。'),
('fig_22_q2_quality','质量响应及ND组合留出','图中展示质量增量与损失变化的拟合及留出关系。ND组合留出RMSE由无质量项的0.207307降至0.062225，改善限于附件B质量情景；模型识别的是θ_Q=λs_Q，λ和s_Q仍需刻度标定。'),
('fig_23_q2_mixture_transfer','配比模型的配方与尺度迁移','图中分别比较配方分组留出和整个尺度留出误差。配方留出RMSE由0.237195降至0.094679，但整尺度迁移明显变差；因此结论适用于已覆盖条件，不能承诺任意规模外推。'),
('fig_24_q2_domain_transfer','领域组合的损失迁移矩阵','矩阵给出配方域变动对验证域预测损失的响应。非对角效应体现跨域迁移，受配方约束与对数链接影响，不应解释为直接估计的领域因果交互。')]: add(Q2,*row)

Q3='问题三 算力约束下的条件资源优化'
for row in [
('fig_25_q3_budget_allocation','预算变化下的最优资源配置','该图展示2048上下文下随预算变化的参数量、训练数据量、预测损失和质量投入。主情景在研究预算范围内持续达到样本参照质量上限，表示该条件设定下上限活跃，不是实测训练规律。'),
('fig_26_q3_cost_shares','训练、注意力与提质预算份额','堆叠柱比较三类预算在不同总算力下的占比。注意力/训练成本比由上下文长度设定，不是独立优化变量；预算构成依赖题设成本代理。'),
('fig_27_q3_context_sensitivity','上下文窗口与成本函数敏感性','曲线比较不同上下文长度和成本函数下的最优量及损失变化。上下文只作为成本情景输入，结果不表示所有训练过程都采用这些窗口。'),
('fig_28_q3_quality_transitions','质量状态转移检查','图中检验预算变化时R0、R1、R2状态是否切换。主情景没有检测到转移并持续处于R2；不能据此把主结果改写为三阶段。'),
('fig_29_q3_sensitivity_gain','不同质量上限的损失收益敏感性','该图并列比较五档筛选参照及条件质量增量带来的预测收益。更高筛选强度支持更大的质量上限，但尾部样本较少，筛选支持不等于可供应训练Token。'),
('fig_30_q3_transition_zoom','敏感性转移点局部放大','放大图定位敏感性情景中状态切换附近的预算区间。已发现的转移只出现在s_Q=2、幂函数成本和2048/4096/8192上下文情景，不代表主情景存在转移。'),
('fig_31_q3_screening_evidence','筛选上限的数据支持','图中展示域内高分筛选与参照均值的关系，用于确定80%、50%、20%、10%、5%并列上限情景。质量上限来自现有评分样本，不是天然质量极限。'),
('fig_32_q3_global_decision_phase','全局质量收益阈值与决策区间','图中展示开始提质阈值λ_on和达到上限阈值λ_full所界定的全局决策区间。阈值来自内层重新优化后的割线比较，局部导数不能替代全局阈值。'),
('fig_33_q3_cap_sensitivity','筛选比例与上限敏感性','图中对照五档质量上限、成本场景与上下文设定。所有情景都是条件预测；book域5%参照仅9篇，极端尾部结果需谨慎解读。'),
('fig_34_q3_resource_tradeoffs','资源影子价值与等效算力','图中展示预算增加时的边际价值、提质投入及相对无提质的等效算力。它只在当前成本、接口和质量刻度假设下成立，不是相对参考版的实测优势。'),
('fig_35_q3_budget_composition','预算分配组成','三元图展示参数训练、注意力和提质预算的组成变化。注意力份额受固定窗口决定，因此图中三元结构不能被解释为三个彼此独立的决策维度。'),
('fig_36_q3_critical_benefit_zoom','临界质量收益局部核验','图中放大关键质量收益区间，核对阈值附近的状态和预算响应。主结论以完整全局比较为准，局部图只辅助检查临界区域。')]: add(Q3,*row)

Q4='问题四 技术演进中介分解与前沿预测'
for row in [
('fig_37_q4_causal_structure','规模中介与非规模路径假设','结构图明确时间经ln N、ln D影响规模响应，并保留剩余时间路径作为非规模进步的解释假设。该路径依赖模型识别条件，不能单凭图示称为因果已证实。'),
('fig_38_q4_data_and_evolution','样本构成与能力演进','图中结合完整N/D基础模型、榜单月度分位和来源审计展示数据覆盖与时间变化。历史归因主样本为35个基础模型、10家机构；不同来源的分数不混合拟合。'),
('fig_39_q4_mediation_and_sensitivity','规模与非规模贡献及混杂敏感性','历史对比中规模贡献为2.939分（56.98%），非规模贡献为2.219分（43.02%）。技术贡献95%机构自助区间包含零；右侧压力测试说明结果会随未观测混杂假设变化。'),
('fig_40_q4_frontier_forecasts','未来能力前沿条件预测','图中对照Q3耦合基础模型前沿与N-only描述性预测。预测起点为2025-03-13，12/24个月结果依赖预算增长与技术半衰期假设，不是从当前日期起算的实时预测。'),
('fig_41_q4_loss_benchmark_bridge','Loss到基准能力的分层桥接','图中分层展示高可比和中可比样本的Loss—能力关系及目标Loss区间。高可比桥接不支持目标Loss时返回缺失；中可比只作为跨报告情景。'),
('fig_42_q4_task_level_audit','逐任务评测与MuSR审计','图中展示三个MuSR子任务归一化聚合和重建结果审计。损坏JSON保留原件及台账，无法解析的目录退出逐任务分析，不用补值制造完整样本。'),
('fig_43_q4_validation','结构模型与短期回测验证','图中比较机构留出、发布队列阻塞留出和榜单月度短期回测。只有对话/微调类别优于上月延续；基础与合并模型未胜出，短期回测也不足以证明长期前沿准确。')]: add(Q4,*row)

assert len(items)==43
# Document

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
styles=doc.styles
for nm in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 st=styles[nm];st.font.name='Noto Sans SC';st._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans SC')
styles['Normal'].font.size=Pt(10.5);styles['Normal'].font.color.rgb=RGBColor(45,52,58)
styles['Normal'].paragraph_format.space_after=Pt(5);styles['Normal'].paragraph_format.line_spacing=1.18
styles['Title'].font.name='Noto Sans SC';styles['Title']._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans SC');
# Remove Word's built-in title paragraph border, which otherwise renders as an unwanted rule.
title_ppr=styles['Title']._element.get_or_add_pPr()
for border in title_ppr.findall(qn('w:pBdr')): title_ppr.remove(border)
styles['Title'].font.size=Pt(25);styles['Title'].font.bold=True;styles['Title'].font.color.rgb=RGBColor(28,57,82)
styles['Heading 1'].font.name='Noto Sans SC';styles['Heading 1']._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans SC');styles['Heading 1'].font.size=Pt(18);styles['Heading 1'].font.color.rgb=RGBColor(28,72,100)
styles['Heading 2'].font.name='Noto Sans SC';styles['Heading 2']._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans SC');styles['Heading 2'].font.size=Pt(13);styles['Heading 2'].font.color.rgb=RGBColor(42,96,122)
styles['Heading 3'].font.name='Noto Sans SC';styles['Heading 3']._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans SC');styles['Heading 3'].font.size=Pt(11);styles['Heading 3'].font.color.rgb=RGBColor(50,70,84)

p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('四问结果分析与论文插图')
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('按 Notebook 执行顺序整理的 43 张 SVG 矢量图');r.font.size=Pt(11);r.font.color.rgb=RGBColor(100,112,122)
doc.add_paragraph('本册将四问正式 Notebook 的计算结果按执行步骤归档。每幅图前先说明所检验的关系和可支持的结论，再嵌入对应 SVG；训练拟合、留出验证、条件预测与未识别假设分别说明。问题一至四各自独立成章，图号按正文预计出现顺序连续编号。')

lastq=None; prev=None
for q,fname,heading,explain in items:
 if q!=lastq:
  if lastq is not None: doc.add_page_break()
  doc.add_heading(q,level=1);lastq=q;prev=None
 stem=fname+'.svg'; svg=FIG/stem; png=FIG/(fname+'.png')
 assert svg.exists() and png.exists(),fname
 # derive execution step from sequential figure within section
 if q==Q1: step=int(fname[4:6])
 elif q==Q2: step=int(fname[4:6])-20
 elif q==Q3: step=int(fname[4:6])-24
 else: step=int(fname[4:6])-36
 h=doc.add_heading(f'执行步骤 {step}  {heading}',level=2)
 h.paragraph_format.keep_with_next=True
 par=doc.add_paragraph(explain);par.paragraph_format.keep_with_next=True
 # preserve vector in Office; 300 dpi PNG is retained as compatibility fallback.
 pic=doc.add_picture(str(png),width=Inches(6.12))
 pic_p=pic._inline.docPr
 pic_p.set('descr',heading+'。'+explain)
 pic_p.set('title',heading)
 pic._inline.graphic.graphicData.pic.nvPicPr.cNvPr.set('name',fname)
 # The next pass will attach the SVG alternative to this PNG-backed drawing.
 cap=doc.add_paragraph(f'图 {int(fname[4:6])}  {heading}')
 cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.space_after=Pt(10)
 for rr in cap.runs: rr.font.size=Pt(9);rr.font.color.rgb=RGBColor(105,112,118)

doc.core_properties.title='四问结果分析与论文插图'
doc.core_properties.subject='问题一至问题四执行步骤对应结果说明及SVG图'
doc.core_properties.author=''
doc.save(OUT)

# Add editable SVG vector alternatives to each embedded PNG drawing using Office DrawingML.
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships','asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main','rel':'http://schemas.openxmlformats.org/package/2006/relationships','ct':'http://schemas.openxmlformats.org/package/2006/content-types'}
qroot=etree.QName
with ZipFile(OUT,'r') as zin:
 blobs={n:zin.read(n) for n in zin.namelist()}
xml=etree.fromstring(blobs['word/document.xml']);rels=etree.fromstring(blobs['word/_rels/document.xml.rels']);cts=etree.fromstring(blobs['[Content_Types].xml'])
# map drawing inline order to figures in document order
blips=xml.xpath('.//a:blip',namespaces=NS)
assert len(blips)==43,len(blips)
for i,(blip,item) in enumerate(zip(blips,items),1):
 q,fname,_,_=item; svg_path=FIG/(fname+'.svg'); media_name=f'vector_{i:02d}.svg'
 blobs['word/media/'+media_name]=svg_path.read_bytes()
 rid=f'rIdSVG{i:03d}'
 rel=etree.SubElement(rels,qroot(NS['rel'],'Relationship'));rel.set('Id',rid);rel.set('Type','http://schemas.openxmlformats.org/officeDocument/2006/relationships/image');rel.set('Target','media/'+media_name)
 extlst=blip.find(qroot(NS['a'],'extLst'))
 if extlst is None: extlst=etree.SubElement(blip,qroot(NS['a'],'extLst'))
 ext=etree.SubElement(extlst,qroot(NS['a'],'ext'));ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
 svgblip=etree.SubElement(ext,qroot(NS['asvg'],'svgBlip'));svgblip.set(qroot(NS['r'],'embed'),rid)
if not cts.xpath('./ct:Default[@Extension="svg"]',namespaces=NS):
 el=etree.SubElement(cts,qroot(NS['ct'],'Default'));el.set('Extension','svg');el.set('ContentType','image/svg+xml')
blobs['word/document.xml']=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
blobs['word/_rels/document.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
blobs['[Content_Types].xml']=etree.tostring(cts,xml_declaration=True,encoding='UTF-8',standalone=True)
tmp=OUT.with_suffix('.tmp.docx')
with ZipFile(tmp,'w',ZIP_DEFLATED) as zout:
 for n,data in blobs.items():zout.writestr(n,data)
tmp.replace(OUT)
print(OUT)

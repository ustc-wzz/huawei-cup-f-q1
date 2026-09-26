"""Edit Q1 sections 5.6–5.10 only, retaining native equations and other chapters.

Usage: bundled-python Q1/restructure_paper_56_510.py SOURCE.docx [OUTPUT.docx]
Existing results are read without fitting or changing any computational model.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from collections import Counter
import csv
import hashlib
import json
import sys
from lxml import etree as E

ROOT = Path(__file__).resolve().parents[1]
SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / 'F题论文初稿_问题一5.6至5.10精简_20260926.docx'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
def tag(s):
    pre, name = s.split(':')
    return '{' + NS[pre] + '}' + name
def text(n):
    return ''.join(n.xpath('.//w:t/text() | .//m:t/text()', namespaces=NS))
def canonical(n):
    return E.tostring(n, method='c14n')
def sub(n, name, **kw):
    item = E.SubElement(n, tag(name))
    for k, v in kw.items(): item.set(tag('w:' + k), str(v))
    return item
def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

source_hash = sha(SRC)
with ZipFile(SRC) as z:
    infos = z.infolist()
    parts = {i.filename: z.read(i.filename) for i in infos}
r = E.fromstring(parts['word/document.xml'])
b = r.find('w:body', NS)
old = list(b)
start = next(i for i,n in enumerate(old) if text(n) == '单纯形上的对数混料响应模型')
end = next(i for i,n in enumerate(old[start:], start) if text(n) == '求解结果')
segment = old[start:end]
by_text = {text(n): n for n in segment}
display_before = Counter(canonical(n) for n in r.xpath('.//m:oMathPara', namespaces=NS))
used = {}
tab = ROOT / 'output_q1/length_domain_calibrated22/tables'
def rows(name):
    p = tab / (name + '.csv'); used[str(p.relative_to(ROOT))] = sha(p)
    with p.open(encoding='utf-8-sig', newline='') as f: return list(csv.DictReader(f))
def data(name):
    p = tab / (name + '.json'); used[str(p.relative_to(ROOT))] = sha(p)
    return json.loads(p.read_text())
inc = rows('quality_increment')
qonly = rows('quality_only_model')
weights = rows('weight_sensitivity')
metrics = rows('mixture_model_metrics_long')
opt = rows('optimal_mixture')
cross = data('main_mixture_crosscheck')
interface = data('q1_outputs_for_q2_q3')
protocol = data('mixture_selection_protocol')
assert 'KKT' in interface['p_star_note']
assert all(abs(interface['p_star'][k] - interface['p_star_loglinear'][k]) < 1e-12 for k in interface['p_star'])
assert all(float(v['对数特征系数 a_i']) < 0 for v in opt)
assert protocol['main_model'] == '对数线性混料'
def metric(model, dataset, key):
    return float(next(x for x in metrics if x['模型'] == model and x['数据集'] == dataset)[key])
def first(prefix):
    return next(n for n in segment if text(n).startswith(prefix))

def inline_text(parent, s):
    import re
    # Match standalone mathematical identifiers, not parts of dataset domain names.
    pattern=r'(?<![A-Za-z_])([A-Za-zα-ωΑ-Ω])_([A-Za-z0-9]+(?:,[A-Za-z0-9]+)?)(?:\^\(([^)]+)\))?(?![A-Za-z_])'
    cursor=0
    for match in re.finditer(pattern,s):
        if match.start()>cursor:sub(sub(parent,'w:r'),'w:t').text=s[cursor:match.start()]
        math=E.SubElement(parent,tag('m:oMath'))
        if match.group(3):
            outer=E.SubElement(math,tag('m:sSup'));base=E.SubElement(outer,tag('m:e'))
        else:base=math
        small=E.SubElement(base,tag('m:sSub'))
        for name,value in [('e',match.group(1)),('sub',match.group(2))]:
            element=E.SubElement(small,tag('m:'+name));rr=E.SubElement(element,tag('m:r'));E.SubElement(rr,tag('m:t')).text=value
        if match.group(3):
            sup=E.SubElement(outer,tag('m:sup'));rr=E.SubElement(sup,tag('m:r'));E.SubElement(rr,tag('m:t')).text='('+match.group(3)+')'
        cursor=match.end()
    if cursor<len(s):sub(sub(parent,'w:r'),'w:t').text=s[cursor:]

def para(s, level=None, appendix=False, pagebreak=False):
    n = E.Element(tag('w:p')); pp = sub(n,'w:pPr')
    sub(pp,'w:keepLines')
    if level and not appendix:
        sub(pp,'w:pStyle',val=str(level))
        sub(pp,'w:spacing',before=120,after=120)
        sub(pp,'w:keepNext')
    elif appendix and level:
        sub(pp,'w:spacing',before=180,after=120)
        sub(pp,'w:keepNext')
        sub(pp,'w:outlineLvl',val=0 if level==1 else 1)
        sub(sub(pp,'w:numPr'),'w:numId',val=0)
        sub(pp,'w:ind',firstLine=0)
    else:
        sub(pp,'w:widowControl')
    if pagebreak: sub(pp,'w:pageBreakBefore')
    run = sub(n,'w:r')
    if appendix and level:
        rp = sub(run,'w:rPr');sub(rp,'w:b');sub(rp,'w:sz',val=32 if level==1 else 26)
        sub(rp,'w:rFonts',ascii='Times New Roman',hAnsi='Times New Roman',eastAsia='黑体')
    if level:sub(run,'w:t').text=s
    else:
        n.remove(run);inline_text(n,s)
    return n

equation_map = {}
def equation(old_number, new_number):
    n = deepcopy(next(x for x in segment if x.tag==tag('w:tbl') and text(x).endswith(f'(5-{old_number})')))
    p = n.findall('w:tr/w:tc', NS)[-1].find('w:p',NS)
    for child in list(p):
        if child.tag not in {tag('w:pPr'),tag('w:bookmarkStart'),tag('w:bookmarkEnd')}:
            p.remove(child)
    run = E.Element(tag('w:r'));sub(run,'w:t').text='('+new_number+')'
    ends=p.findall('w:bookmarkEnd',NS)
    if ends:p.insert(p.index(ends[0]),run)
    else:p.append(run)
    equation_map[f'5-{old_number}']=new_number
    # Reserve space between the formula and label without changing math content.
    widths=[250,7522,1300]
    for cell,w in zip(n.findall('w:tr/w:tc',NS),widths):
        cell.find('w:tcPr/w:tcW',NS).set(tag('w:w'),str(w))
    for col,w in zip(n.findall('w:tblGrid/w:gridCol',NS),widths):col.set(tag('w:w'),str(w))
    return n

# Evidence is read from frozen outputs; these calculations only summarize existing tables.
import numpy as np
marginal = rows('domain_marginal_effects')
transfer = rows('transfer_matrix')
mr = {x['']:x for x in marginal}
val_domains = list(transfer[0])[1:]
train_domains = [x[''] for x in transfer]
pm = np.array([float(mr[d]['训练集平均份额']) for d in train_domains])
T = np.array([[float(x[v]) for v in val_domains] for x in transfer])
G = T/(pm[:,None]+.001)
G -= (G*pm[:,None]).sum(axis=0)
same_ranks = {v:int(np.argsort(G[:,j]).tolist().index(train_domains.index(v))+1) for j,v in enumerate(val_domains)}
assert sum(v==1 for v in same_ranks.values())==7
assert abs(sum(interface['p_star'].values())-1)<1e-12 and min(interface['p_star'].values())>=0
used['notebook_source.py']=sha(ROOT/'notebook_source.py')
# Fail rather than silently retain rounded claims if the source results change.
def claim(value, expected):
    assert f'{float(value):.4f}'==expected,(value,expected)
for row,before,after in zip(inc,['0.9316','0.9116','0.7548'],['0.9327','0.9127','0.7478']):
    claim(row['对数线性(p) Spearman'],before);claim(row['对数线性(p)+Qbar Spearman'],after)
claim(next(x['R2(仅Qbar)'] for x in qonly if '1M 检验' in x['数据集']),'-0.0480')
for key,expected in [('GBM_loss_main','4.7441'),('GBM_loss_train_mean','4.7638'),('GBM_loss_uniform','4.7443'),('GBM_loss_search_best','4.7052')]:claim(cross[key],expected)
claim(mr['pile_cc']['边际效应@训练均值配比'],'0.0507');claim(mr['pile_cc']['边际效应@均匀配比'],'-0.3498')
single=[x for x in weights if '权重×' in x['方案']]
assert len(single)==44 and all(abs(float(x['域均值Spearman'])-1)<1e-12 for x in single)
claim(min(float(x['样本Spearman']) for x in single),'0.9990')
equal=next(x for x in weights if '22指标等权' in x['方案'])
claim(equal['样本Spearman'],'0.9465');claim(equal['域均值Spearman'],'0.7857')
for ds,rmse,rho in [('1M 检验集 (A6/A7)','0.1071','0.9316'),('60M 检验集 (A8/A9)','1.4894','0.9116'),('1B 检验集 (A10/A11)','2.6679','0.7548')]:
    claim(metric('对数线性混料',ds,'RMSE'),rmse);claim(metric('对数线性混料',ds,'Spearman'),rho)
claim(metric('对数线性混料','1M 检验集 (A6/A7)','R2'),'0.8514')
mapped=['arxiv','wikipedia_en','github','stackexchange','gutenberg_pg_19','pile_cc']
qa=np.array([interface['log_linear_model']['a_i'][d] for d in mapped])
qq=np.array([interface['mixture_domain_Q'][d] for d in mapped])
rank=lambda x:np.array([np.where(np.sort(x)==v)[0].mean()+1 for v in x])
claim(np.corrcoef(qa,qq)[0,1],'-0.4368');claim(np.corrcoef(rank(qa),rank(qq))[0,1],'-0.4857')
for key,pct in [('stackexchange','12.15'),('dm_mathematics','11.97'),('pile_cc','10.44')]:assert f'{interface["p_star"][key]*100:.2f}'==pct


# Compact editable three-line tables, using the source text family and page width.
def simple_table(headers, records, widths, size=22):
    t=E.Element(tag('w:tbl'));pr=sub(t,'w:tblPr')
    sub(pr,'w:tblW',w=sum(widths),type='dxa');sub(pr,'w:tblLayout',type='fixed')
    borders=sub(pr,'w:tblBorders')
    for edge in ['top','bottom']:sub(borders,'w:'+edge,val='single',sz=8,color='000000')
    for edge in ['left','right','insideH','insideV']:sub(borders,'w:'+edge,val='nil')
    grid=sub(t,'w:tblGrid')
    for width in widths:sub(grid,'w:gridCol',w=width)
    for ri,record in enumerate([headers]+records):
        row=sub(t,'w:tr');rpr=sub(row,'w:trPr');sub(rpr,'w:cantSplit')
        if ri==0:sub(rpr,'w:tblHeader')
        for value,width in zip(record,widths):
            cell=sub(row,'w:tc');cp=sub(cell,'w:tcPr');sub(cp,'w:tcW',w=width,type='dxa')
            if ri==0:sub(sub(cp,'w:tcBorders'),'w:bottom',val='single',sz=6,color='000000')
            p0=sub(cell,'w:p');pp=sub(p0,'w:pPr');sub(pp,'w:ind',firstLine=0,firstLineChars=0,left=0,right=0,leftChars=0,rightChars=0)
            sub(pp,'w:jc',val='left')
            if ri<len(records):sub(pp,'w:keepNext')
            sub(pp,'w:spacing',before=0,after=0,line=276,lineRule='auto')
            rr=sub(p0,'w:r');rp=sub(rr,'w:rPr');sub(rp,'w:rFonts',ascii='Times New Roman',hAnsi='Times New Roman',eastAsia='宋体');sub(rp,'w:sz',val=size)
            if ri==0:sub(rp,'w:b')
            sub(rr,'w:t').text=str(value)
    return t

main=[]
def p(s,level=None):main.append(para(s,level))
def cp(prefix):main.append(deepcopy(first(prefix)))
def eq(a,c):main.append(equation(a,c))
p('单纯形上的对数混料响应模型',2)
p('以17个训练领域的份额向量p为输入，以同一训练尺度下13个验证域的平均损失L(p)为预测目标。份额满足p_i≥0且总和为1，构成单纯形Δ¹⁶。采用对数加性响应描述配比与损失的关系：')
eq(19,'5-18')
p('其中c为截距，a_i为领域i的对数特征系数，平滑常数ε_p=10⁻³使零份额处仍有有限预测值，并影响低份额区域的曲率。该模型服务于固定训练条件下的配比分析；完整规模响应由问题二处理。')
p('对M=512个训练配方，记响应向量为y，设计矩阵H的截距列为1，其余列为H_mi=ln(p_mi+ε_p)，以最小二乘估计参数：')
eq(20,'5-19')
p('拟合时不预先限制系数符号，平滑常数按5.10节的训练内验证程序选择。局部近似依据、设计矩阵的可识别性及系数重抽样区间见附录一。')
p('解释配比调整时，必须同时计入被替代领域的损失变化。取p(t)=(1−t)p+t e_i，其中e_i为第i个单位向量、0≤t≤1，表示按原比例减少其他领域并增加领域i，其起点方向效应为：')
eq(23,'5-20')
p('g_i(p)<0表示沿该可行方向作微小调整可降低预测损失。以pile_cc为例，g_i在训练平均配比处为0.0507，在均匀配比处为−0.3498，说明调整收益依赖起始配比，不能直接按|a_i|作统一排序。完整系数、区间及局部替代图见附录一辅助分析。')

p('模型比较与逐验证域响应',2)
p('四类模型均使用相同配方与平均损失目标。表5-1给出结构差异及训练内外层五折的折外预测结果；均方根误差（RMSE）越小，预测误差越低。具体划分与参数选择见5.10节。')
p('表5-1 配比响应模型比较')
models=['线性混料','对数线性混料','加性样条（PDF）','梯度提升树 GBM']
main.append(simple_table(['模型','响应结构与用途','折外RMSE'],[
 ['线性混料','无截距的份额线性加和；线性基线',f'{metric(models[0],"外层5折CV","RMSE"):.4f}'],
 ['对数混料','平滑对数加和；解释与约束求解主模型',f'{metric(models[1],"外层5折CV","RMSE"):.4f}'],
 ['加性样条','各份额的平滑曲线加和，无显式交互',f'{metric(models[2],"外层5折CV","RMSE"):.4f}'],
 ['GBM','梯度提升回归树；刻画非线性与交互',f'{metric(models[3],"外层5折CV","RMSE"):.4f}']], [1900,5672,1500]))
p('GBM的折外误差最低。对数模型保留了明确的可行调整方向，并能在满足条件时求得约束全局最优配比，故选为解释与优化的主模型，GBM用于预测比较和候选核查；这一选择不意味着对数模型的预测精度全面最优。')
p('为区分不同验证任务，对每个验证域v单独拟合相同形式的响应：')
eq(26,'5-21')
p('其中L_v为验证域v的损失，c_v与a_i^(v)为对应回归参数。T保存17×13个对数特征系数；在训练平均配比处，将各列系数代入式（5-20），得到可行方向效应矩阵G，负值表示该方向局部降低对应验证域损失。')
p('按G的列内效应由低到高排序，13个验证域中有7个以同名训练域最有利；pile_cc与pubmed_abstracts的同名域分别排第4和第3。领域匹配并非唯一决定因素，这些结果只描述给定配比附近的模型响应。完整T、G热图及对照参数见附录一。')

p('质量指标的增量解释检验',2)
cp('质量映射采用文档均分：')
eq(27,'5-22')
p('直接、近似直接与推断映射的κ_i分别取1、0.8、0.5，以控制向总体均值收缩的程度。17个配方域中有6个具有参考质量映射，其余11个以μ_Q回填；回填值属于参考设定。固定领域质量向量q后，定义配比加权质量：')
eq(28,'5-23')
p('P为各实验配方组成的份额矩阵，Q_mix为对应质量指数向量。仅质量线性模型在1M检验集上的决定系数为−0.0480，表明这一标量单独解释损失的能力有限。进一步将其加入对数主模型：')
eq(32,'5-24')
p('其中c_p,Q、a_i^(p,Q)与θ_p,Q均由相同训练集估计。对1M、60M、1B检验集，加入质量项前后的Spearman秩相关系数（衡量配方损失排序一致性）分别为0.9316→0.9327、0.9116→0.9127、0.7548→0.7478，未显示跨尺度一致的预测增益。因此候选求解继续使用不含质量项的对数主模型。')
p('领域质量固定时，配比加权质量完全由配比决定，因此本节检验统计关联与预测增益，不能据此识别独立质量效应。两类不可识别性推导、仅质量基线及质量与系数的相关分析见附录一。')

p('基于主模型的候选配比求解',2)
p('固定训练尺度与评分条件，在份额非负且总和为1的约束下，最小化对数主模型的预测损失：')
eq(33,'5-25')
p('当所有a_i≤0且至少一个严格为负时，目标为凸函数；当前17个拟合系数均为负，目标严格凸。由约束最优性条件得到候选解：')
eq(36,'5-26')
p('其中λ>0为份额和约束的乘子，由份额和等于1确定。通过二分法求解，并用序列最小二乘约束优化（SLSQP）交叉核查，得到本模型在上述约束下的唯一全局最优解；系数不满足条件时须另行求解。完整推导与数值设置见附录一。')
p('该解为本文主配方，份额非负且总和为1；前三个领域为stackexchange（12.15%）、dm_mathematics（11.97%）和pile_cc（10.44%）。GBM对主配方、训练平均配比和均匀配比的预测损失分别为4.7441、4.7638和4.7443，主配方优于训练平均配比，与均匀配比差异很小。GBM搜索与两模型名次融合只提供辅助候选。')
p('候选在单纯形内可行，但这不足以证明其处于训练配方的联合支持范围；逐域坐标范围核查也不能替代联合支持检查。因此这里得到的是代理模型的预测最优配比，真实收益仍需新增训练实验验证。')

p('模型验证与稳健性分析',2)
p('质量评分稳健性',3)
p('固定A1参照，分别将单项权重上下扰动20%后归一化，并比较去冗余强度0.5、1、1.5、2及等权方案；短板参数β取0、0.1、0.25、0.5、0.75、1，长度校准比较10、20、30组。以样本排序、前列样本重合程度和领域均值变化评价对设置的依赖，主设置保持β=0.25。')
p('44项单权重扰动中，样本排序与主方案的Spearman相关系数最低为0.9990，领域均值排序均不变；改用等权方案时，样本与领域均值秩相关分别为0.9465和0.7857。评分对局部权重扰动较稳定，但赋权规则的整体改变仍会影响领域排序。')
p('稳定性不直接证明评分准确反映训练价值。长度留出检验评价长度校准组件，机械重复与格式扰动评价指标响应；完整设置与核验流程见附录一复现设置。')
p('配比预测与跨尺度检验',3)
p('在A4/A5的512个训练配方内进行嵌套五折验证：外层比较模型，有待选参数的模型在各外层训练折内再用内层五折选择；每个配方及其13域损失作为同一划分单位。随后在全训练集重拟合并冻结参数，再用于A6—A11检验。各模型的折外误差见表5-1。')
p('1M与60M检验集各含256个配方，1B检验集含64个配方。对数主模型在同尺度1M检验集的RMSE为0.1071、决定系数为0.8514；冻结后在60M、1B上的秩相关分别为0.9116、0.7548，而RMSE增至1.4894、2.6679。跨尺度排序保持不等于绝对损失预测准确。A12—A15估算或外推表仅用于一致性分析，不作为新增独立实测验证。')
p('前期分析曾查看固定检验集的平滑参数表现；本轮虽在A4/A5内选参，A6—A11仍属于已被查看过的固定检验集，不能表述为全新未接触的外部验证。各模型跨尺度散点、评价指标定义与复现流程见附录一。')

app=[]
def ap(s,level=None,pagebreak=False):app.append(para(s,level,appendix=True,pagebreak=pagebreak))
def ac(prefix):app.append(deepcopy(first(prefix)))
def ae(a,c):app.append(equation(a,c))
ap('附录一 问题一补充推导与复现设置',1,True)
ap('一 补充推导',2)
ap('对数响应的局部近似与参数解释。固定总训练量D，若领域数据不足项近似为B_i乘以(D p_i)的−ν_i次幂，令t_i=ln p_i，在内部参考配比p₀附近作一阶展开：')
ae(18,'附1-1')
ap('其中B_i、ν_i为领域不足项的幅度与指数。将常数合并可得到对数加性近似，但其成立依赖可分解假设和局部范围；正文平滑模型不据此声称在整个单纯形上准确。由正文式（5-18）得到：')
ae(21,'附1-2')
ap('仅在a_i<0时，单项响应才呈随份额增加而减弱的损失改善。a_i为对数特征系数，其精确关系为：')
ae(22,'附1-3')
ap('损失的无量纲弹性还需将相应导数除以预测损失；上述非约束偏导也不能直接替代正文式（5-20）的可行方向效应。')
ap('设计矩阵与可识别性。正文最小二乘参数的唯一识别要求H具有足够独立变化。原始份额满足总和为1，故线性混料基准不另设截距；对数特征不自动保证消除共线性。预测误差与参数稳定性应分别检查。')
ap('固定质量下的识别限制。若在线性份额模型中增加配比加权质量项，有：')
ae(29,'附1-4')
ap('θ_Q可被各领域系数吸收。对未平滑的可分对数有效数据量形式，设φ(q_i)>0，正份额下有：')
ae(30,'附1-5')
ap('固定D与q后，后两项均为常数，可被截距吸收。该推导针对所写参数化，不推广到任意响应模型；独立质量干预仍需在相同配比下改变质量并观测损失。')
ap('约束最优性。正文式（5-25）的拉格朗日函数为：')
ae(34,'附1-6')
ap('λ为等式约束乘子，μ_i为非负约束乘子。KKT条件＝一阶驻点、原始与对偶可行性及互补条件；除p_i≥0、份额和为1、μ_i≥0、μ_i p_i=0外，还包括：')
ae(35,'附1-7')
ap('在正文凸性条件下由此得到式（5-26），所有系数严格为负时目标严格凸，最优解唯一。λ的数值区间与交叉核查设置列于第三部分。')

ap('二 辅助分析',2)
ap('系数区间与参考配比。对训练配方及其响应成对重抽样500次，逐次重新拟合，以系数2.5%与97.5%经验分位形成区间。训练平均配比及均匀配比定义为：')
ae(24,'附1-8')
ap('区间条件于既定平滑常数与模型形式；重抽样设计退化时应记录。图附1-1同时给出系数区间及两种配比处的可行方向效应。')
app.append(deepcopy(old[329]));ap('图附1-1 对数系数区间与参考配比处的方向效应（原图12）')
ap('线性基准与局部替代。无截距的一阶线性混料模型为：')
ae(25,'附1-9')
ap('有限替代以训练平均配比为起点：两两转移0.001份额（0.1个百分点），以及从uspto_backgrounds向其他域转移0.01份额（1个百分点）。图附1-2比较三种模型的预测损失变化，仅描述指定方向和幅度的替代，不据此认定真实互补机制。')
app.append(deepcopy(old[334]));ap('图附1-2 固定配比约束下的局部替代效应（原图13）')
ap('逐验证域完整系数。表附1-1与表附1-2给出T的全部17×13个系数，行编号i按下列配方域顺序；数值保留四位小数，局部效应与排序使用未舍入值计算。')
ap('；'.join(f'{i+1}={d}' for i,d in enumerate(train_domains))+'。')
for block,cols in enumerate([val_domains[:7],val_domains[7:]],1):
    ap(f'表附1-{block} 逐验证域对数特征系数（第{1 if block==1 else 8}—{7 if block==1 else 13}列）')
    ap('；'.join(f'v{val_domains.index(c)+1}={c}' for c in cols)+'。')
    app.append(simple_table(['i']+[f'v{val_domains.index(c)+1}' for c in cols], [[str(i+1)]+[f'{float(row[c]):.4f}' for c in cols] for i,row in enumerate(transfer)], [500]+[int(8572/len(cols))]*len(cols),size=21))
ap('将T各列代入正文式（5-20），得到图附1-3的训练平均配比处可行方向效应G。颜色表示局部方向导数，矩阵T本身不等于G。')
# Add the existing, verified figure without redrawing or altering source images.
from docx import Document
from docx.shared import Inches
fig=ROOT/'figures/fig_14_transfer_matrix.png';used[str(fig.relative_to(ROOT))]=sha(fig)
tmpdoc=Document();pic=tmpdoc.add_paragraph();pic.add_run().add_picture(str(fig),width=Inches(5.9))
imgp=E.fromstring(E.tostring(pic._p))
A='http://schemas.openxmlformats.org/drawingml/2006/main';R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
newid='rIdQ1ExpandedHeatmap'
for x in imgp.iter('{'+A+'}blip'):x.set('{'+R+'}embed',newid)
WP='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
maxid=max(int(x.get('id')) for x in r.iter('{'+WP+'}docPr'))
for x in imgp.iter('{'+WP+'}docPr'):x.set('id',str(maxid+1));x.set('name','Q1 appendix direction heatmap')
app.append(imgp);ap('图附1-3 训练平均配比处的17×13可行方向效应')
relroot=E.fromstring(parts['word/_rels/document.xml.rels'])
E.SubElement(relroot,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id=newid,Type=R+'/image',Target='media/q1_expanded_transfer_matrix.png')
parts['word/_rels/document.xml.rels']=E.tostring(relroot,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['word/media/q1_expanded_transfer_matrix.png']=fig.read_bytes()
ap('质量关联基线。仅质量模型为：');ae(31,'附1-10')
ap('质量与系数的相关分析只纳入6个实际映射域，其q_i与a_i的Pearson相关为−0.4368、Spearman秩相关为−0.4857；11个均值回填域不作为独立质量测量点。结果用于描述给定映射下的关联，不承担独立质量效应识别。')
ap('辅助候选。以各分量均为0.7的Dirichlet分布生成50,000个配比，加入训练配方、主模型最优解、训练平均配比及均匀配比，分别按对数模型与GBM预测损失升序排名，定义：');ae(37,'附1-11')
ap('R_log、R_GBM表示各候选的预测名次。随机候选直接满足单纯形约束，不做逐分量截断；名次融合只作启发式对照，主配方仍为正文式（5-26）的解。GBM搜索最佳候选的GBM预测损失为4.7052，低于对数主候选的4.7441，说明两代理模型的最优候选并不一致。')
ap('跨尺度预测散点。图附1-4比较四类模型在1M、60M、1B检验集的预测与实测损失；红色虚线表示预测与实测数值一致的位置，跨尺度散点的整体偏移反映绝对损失尚未校准。')
app.append(deepcopy(old[344]));ap('图附1-4 四类模型的跨尺度检验散点（原图11）')

ap('三 复现设置',2)
ap('训练内参数选择。外层五折随机种子为0，内层为1。对数平滑常数候选为10⁻⁴、10⁻³、10⁻²、0.03；加性样条使用4个均匀节点、二次基函数、线性外推，删除uspto_backgrounds参考域后对特征标准化，岭参数候选为0.001、0.01、0.1、1、10、100、1000。GBM候选为300/500棵树、深度2/3，学习率0.03、子采样率0.8。全训练集最终参数为ε_p=0.001、岭参数10、GBM深度3及500棵树。')
ap('数值求解。λ初始区间为[10⁻⁶,10³]，以几何中点二分200次；SLSQP最大迭代500次、目标容差10⁻¹²，检查优化器成功及两种配比最大分量差小于10⁻⁵，并检查份额非负、份额和及目标值。逐域训练范围属于坐标核查，不代表联合支持。')
ac('评分模块先检查')
ap('同口径误差采用均方根误差与决定系数：');ae(38,'附1-12')
ap('排序采用Spearman秩相关系数，并列值以中秩处理：');ae(39,'附1-13')
ap('若响应或预测近常数，相关系数可能无定义，应保留缺失标记。跨尺度散点的事后拟合直线不能直接充当独立标定，须另设标定样本后评价绝对误差。')
ac('对 A12');ac('各尺度表可分别拟合')
ap('表附1-3 主设置、参数来源与核验方式')
param_table=deepcopy(next(n for n in segment if n.tag==tag('w:tbl') and text(n).startswith('模块参数或规则')))
def replace_cell(cell,s):
    pp=cell.find('w:p',NS)
    for child in list(pp):
        if child.tag!=tag('w:pPr'):pp.remove(child)
    sub(sub(pp,'w:r'),'w:t').text=s
    for extra in cell.findall('w:p',NS)[1:]:cell.remove(extra)
for row in param_table.findall('w:tr',NS):
    cells=row.findall('w:tc',NS)
    if text(cells[0])=='候选搜索':
        replace_cell(cells[2],'50,000；各分量0.7');replace_cell(cells[3],'可行性及超训练坐标范围检查')
    if text(cells[0])=='对数混料':replace_cell(cells[3],'训练内选参；保留历史检验查看说明')
    if text(cells[0])=='冲突识别':replace_cell(cells[2],'r≤−0.30；20%/80%')
    if text(cells[0])=='GBM':replace_cell(cells[3],'外层训练折内选参；同划分折外比较')
app.append(param_table)
ap('表附1-4 问题一计算流程')
flow=deepcopy(next(n for n in segment if n.tag==tag('w:tbl') and text(n).startswith('步骤操作')))
for row in flow.findall('w:tr',NS):
    cells=row.findall('w:tc',NS)
    if text(cells[0])=='10':replace_cell(cells[1],'可行方向效应、有限替代及逐验证域拟合')
    if text(cells[0])=='12':replace_cell(cells[2],'单纯形可行候选与训练支持边界')
app.append(flow);ac('对后续质量成本')

# Compact the two appendix tables while keeping the inherited three-line style.
for table in (param_table, flow):
    for paragraph in table.xpath('.//w:tc/w:p', namespaces=NS):
        pp=paragraph.find('w:pPr', NS)
        if pp is None:
            pp=E.Element(tag('w:pPr'));paragraph.insert(0,pp)
        spacing=pp.find('w:spacing', NS)
        if spacing is None: spacing=sub(pp,'w:spacing')
        for k,v in {'before':'0','after':'0','line':'276','lineRule':'auto'}.items():
            spacing.set(tag('w:'+k),v)
        for run in paragraph.findall('w:r', NS):
            rp=run.find('w:rPr', NS)
            if rp is None: rp=E.Element(tag('w:rPr'));run.insert(0,rp)
            size=rp.find('w:sz', NS)
            if size is None:size=sub(rp,'w:sz')
            size.set(tag('w:val'),'22')

# Deliberate table pagination and caption formatting.
for i,n in enumerate(app):
    if n.tag==tag('w:p') and text(n).startswith(('表附1-', '图附1-')):
        pp=n.find('w:pPr',NS)
        sub(pp,'w:ind',firstLine=0,firstLineChars=0)
        sub(pp,'w:jc',val='center')
        if text(n).startswith('表附1-'):
            sub(pp,'w:keepNext')
        if text(n).startswith(('表附1-3 ', '表附1-4 ')):
            sub(pp,'w:pageBreakBefore')
# Keep image and caption together.
for previous,current in zip(app,app[1:]):
    if previous.xpath('.//w:drawing',namespaces=NS):
        pp=previous.find('w:pPr',NS)
        if pp is None:pp=E.Element(tag('w:pPr'));previous.insert(0,pp)
        if pp.find('w:keepNext',NS) is None:sub(pp,'w:keepNext')
# Keep an equation/table with its immediately preceding introduction or caption.
for group in (main, app):
    for previous, current in zip(group, group[1:]):
        if current.tag == tag('w:tbl') and previous.tag == tag('w:p'):
            pp = previous.find('w:pPr', NS)
            if pp is None:
                pp = E.Element(tag('w:pPr')); previous.insert(0, pp)
            if pp.find('w:keepNext', NS) is None: sub(pp,'w:keepNext')
            if pp.find('w:keepLines', NS) is None: sub(pp,'w:keepLines')

for n in segment:b.remove(n)
for i,n in enumerate(main):b.insert(start+i,n)
app_index=len(b)-1
assert b[-1].tag==tag('w:sectPr')
for n in app:b.insert(len(b)-1,n)
after=list(b)
assert [canonical(n) for n in after[:start]]==[canonical(n) for n in old[:start]]
tail_start=start+len(main)
assert [canonical(n) for n in after[tail_start:app_index]]==[canonical(n) for n in old[end:-1]]
assert canonical(b[-1])==canonical(old[-1])
assert display_before==Counter(canonical(n) for n in r.xpath('.//m:oMathPara',namespaces=NS))
bookmarks=r.xpath('.//w:bookmarkStart/@w:id',namespaces=NS)
assert len(bookmarks)==len(set(bookmarks)), 'duplicate bookmark ids'
parts['word/document.xml']=E.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
    for info in infos:z.writestr(info,parts[info.filename])
    z.writestr('word/media/q1_expanded_transfer_matrix.png',parts['word/media/q1_expanded_transfer_matrix.png'])
assert sha(SRC)==source_hash
with ZipFile(SRC) as z:
    changed=[n for n in z.namelist() if z.read(n)!=parts[n]]
assert set(changed)=={'word/document.xml','word/_rels/document.xml.rels'}
old_chars=sum(len(text(n)) for n in segment)
new_chars=sum(len(text(n)) for n in main)
report={'version':'v0.9.16','source':SRC.name,'source_sha256':source_hash,
 'output':OUT.name,'output_sha256':sha(OUT),'sources':used,
 'modified_sections':['5.6','5.7','5.8','5.9','5.10'],'other_body_nodes_unchanged':True,
 'all_original_display_equations_preserved':True,'all_original_parts_except_document_and_image_relationships_unchanged':True,
 'source_unchanged':True,'main_candidate':'loglinear KKT optimum; rank fusion is supplementary',
 'corrected_implementation_notes':['Dirichlet parameter 0.71 → 0.7','No component clipping / renormalizing in current candidate generation'],
 'equation_number_map':equation_map,'old_body_characters':old_chars,
 'new_body_characters':new_chars,'body_reduction_percent':round(100*(1-new_chars/old_chars),1),
 'same_domain_local_ranks':same_ranks,'main_display_equations':9,'appendix_display_equations':13,'original_figures_relocated':3,'original_images_unchanged':True,'render_qa':'pending'}
(ROOT/'Q1/论文5.6至5.10精简正文.md').write_text('\n\n'.join(text(n) for n in main),encoding='utf-8')
(ROOT/'Q1/论文问题一5.6至5.10附录.md').write_text('\n\n'.join(text(n) for n in app),encoding='utf-8')
(ROOT/'Q1/论文5.6至5.10精简核验.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))

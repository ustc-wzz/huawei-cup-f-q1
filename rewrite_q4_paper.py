"""Rewrite only chapter 8 of the supplied paper; retain the original DOCX package.
Usage: bundled-python rewrite_q4_paper.py SOURCE.docx [OUTPUT.docx]
No model or result is recomputed. Numerical sources: output_q4_evolution/tables.
"""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import sys, hashlib, json, re
from lxml import etree as E

ROOT=Path(__file__).resolve().parent
SRC=Path(sys.argv[1])
OUT=Path(sys.argv[2]) if len(sys.argv)>2 else ROOT/'F题论文初稿_问题四重写_20260926.docx'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
def tag(n):a,b=n.split(':');return '{'+NS[a]+'}'+b
def elem(n,**attrs):
 e=E.Element(tag(n))
 for k,v in attrs.items():e.set(tag('w:'+k),str(v))
 return e
with ZipFile(SRC) as z: parts={i.filename:z.read(i.filename) for i in z.infolist()};infos=z.infolist()
r=E.fromstring(parts['word/document.xml']);body=r.find('w:body',NS);old=list(body)
assert ''.join(old[866].xpath('.//w:t/text()',namespaces=NS))=='问题四模型建立与求解'
old_hash=hashlib.sha256(SRC.read_bytes()).hexdigest()
new=[];plain=[];eq_no=0;table_no=0

def scrub(e):
 for x in e.iter():
  for k in list(x.attrib):
   if E.QName(k).localname in ['paraId','textId','rsidR','rsidRDefault','rsidP','rsidTr','rsidRPr']:del x.attrib[k]
 for x in list(e.xpath('.//w:bookmarkStart|.//w:bookmarkEnd|.//w:lastRenderedPageBreak',namespaces=NS)):
  x.getparent().remove(x)
 return e

def mr(t):
 e=E.Element(tag('m:r'));pr=E.SubElement(e,tag('m:rPr'));sty=E.SubElement(pr,tag('m:sty'));sty.set(tag('m:val'),'p')
 tt=E.SubElement(e,tag('m:t'));tt.text=t
 return e

def sub(base,idx):
 e=E.Element(tag('m:sSub'));a=E.SubElement(e,tag('m:e'));a.append(mr(base));b=E.SubElement(e,tag('m:sub'));b.append(mr(idx));return e

def supsub(base,idx,sp):
 e=E.Element(tag('m:sSubSup'))
 for name,t in [('e',base),('sub',idx),('sup',sp)]:
  x=E.SubElement(e,tag('m:'+name));x.append(mr(t))
 return e

INLINE={a:sub(b,c) for a,b,c in [
 ('sᵢₖ','s','ik'),('Mᵢ','M','i'),('a₀','a','0'),('aT','a','T'),('εM,ᵢ','ε','M,i'),('ΣM','Σ','M'),
 ('Tᵢ','T','i'),('β₀','β','0'),('γT','γ','T'),('nM','n','M'),('nY','n','Y'),('sx','s','x'),('sY','s','Y'),
 ('ax','a','x'),('bx','b','x'),('C⋆','C','⋆'),('gC','g','C'),('rC','r','C'),('p₀','p','0'),('sQ','s','Q'),
 ('κτ','κ','τ'),('bN,g','b','N,g'),('bT,g','b','T,g'),('Ch','C','h'),('Iℓ','I','ℓ') ]}
for b in ['N','D','u','L']:INLINE[b+'*h']=supsub(b,'h','⋆')
INLINE_PATTERN=re.compile('|'.join(re.escape(t) for t in sorted(INLINE,key=len,reverse=True)))

def append_text(parent,text,runprops=None):
 start=0
 for mat in INLINE_PATTERN.finditer(text):
  if mat.start()>start:
   rr=elem('w:r')
   if runprops is not None:rr.append(deepcopy(runprops))
   tt=elem('w:t');tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');tt.text=text[start:mat.start()];rr.append(tt);parent.append(rr)
  om=E.Element(tag('m:oMath'));om.append(deepcopy(INLINE[mat.group()]));parent.append(om);start=mat.end()
 if start<len(text):
  rr=elem('w:r')
  if runprops is not None:rr.append(deepcopy(runprops))
  tt=elem('w:t');tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');tt.text=text[start:];rr.append(tt);parent.append(rr)

def p(text,role='body'):
 source={'body':867,'h1':866,'h2':868,'h3':869,'caption':874}[role]
 template=old[source];e=elem('w:p')
 pp=template.find('w:pPr',NS)
 if pp is not None:e.append(deepcopy(pp))
 if role in ['body','caption']:
  pp=e.find('w:pPr',NS)
  if pp is None:pp=elem('w:pPr');e.insert(0,pp)
  # Body inherits the source's default paragraph settings; captions stay source styled.
  if role=='caption':
   if pp.find('w:keepNext',NS) is None:pp.append(elem('w:keepNext'))
   np=pp.find('w:numPr',NS)
   if np is not None:pp.remove(np)
   np=elem('w:numPr');np.append(elem('w:numId',val=0));pp.append(np)
  else:
   if pp.find('w:widowControl',NS) is None:pp.append(elem('w:widowControl'))
 pr=template.find('w:r/w:rPr',NS)
 append_text(e,text,pr)
 new.append(e);plain.append(('## ' if role.startswith('h') else '')+text)
 return e

def eq(source):
 global eq_no
 if new and new[-1].tag==tag('w:p'):
  pp=new[-1].find('w:pPr',NS)
  if pp is None:pp=elem('w:pPr');new[-1].insert(0,pp)
  if pp.find('w:keepNext',NS) is None:pp.append(elem('w:keepNext'))
 eq_no+=1;e=scrub(deepcopy(old[source]));assert e.tag==tag('w:tbl')
 cells=e.findall('w:tr/w:tc',NS);assert len(cells)==3
 target=cells[-1]
 for x in list(target):
  if x.tag!=tag('w:tcPr'):target.remove(x)
 pp=elem('w:p');props=elem('w:pPr');props.append(elem('w:pStyle',val='-'));props.append(elem('w:jc',val='right'));pp.append(props)
 rr=elem('w:r');tt=elem('w:t');tt.text=f'(8-{eq_no})';rr.append(tt);pp.append(rr);target.append(pp)
 new.append(e);plain.append(f'[公式8-{eq_no}] '+''.join(e.xpath('.//m:t/text()',namespaces=NS)))
 return eq_no

def estimate_equation():
 eq(935)
 cell=new[-1].findall('w:tr/w:tc',NS)[1]
 om=cell.find('.//m:oMath',NS)
 for x in list(om):om.remove(x)
 acc=E.Element(tag('m:acc'));pr=E.SubElement(acc,tag('m:accPr'));ch=E.SubElement(pr,tag('m:chr'));ch.set(tag('m:val'),'̂');aa=E.SubElement(acc,tag('m:e'));aa.append(mr('Z'));om.append(acc)
 om.append(mr('=−3.541544+1.217559x(N,D)+0.191554T.'))
 plain[-1]='拟合方程：Ẑ=−3.541544+1.217559x(N,D)+0.191554T。'

def table(title,headers,rows,widths):
 global table_no
 table_no+=1;p(f'表8-{table_no}  {title}','caption')
 base=old[1089];e=elem('w:tbl');e.append(deepcopy(base.find('w:tblPr',NS)))
 grid=elem('w:tblGrid')
 for width in widths:grid.append(elem('w:gridCol',w=width))
 e.append(grid)
 for ri,values in enumerate([headers]+rows):
  tr=elem('w:tr');tp=elem('w:trPr');tp.append(elem('w:cantSplit'))
  if ri==0:tp.append(elem('w:tblHeader'))
  tr.append(tp)
  for ci,(val,width) in enumerate(zip(values,widths)):
   tc=elem('w:tc');cp=deepcopy(base.findall('w:tr',NS)[0 if ri==0 else 1].find('w:tc/w:tcPr',NS));cp.find('w:tcW',NS).set(tag('w:w'),str(width));tc.append(cp)
   pp=elem('w:p');pr=elem('w:pPr');pr.append(elem('w:pStyle',val='-'));pr.append(elem('w:jc',val='left' if ci==0 else 'center'));pr.append(elem('w:spacing',before=0,after=0));
   if ri<len(rows):pr.append(elem('w:keepNext'))
   pp.append(pr)
   rr=elem('w:r');rp=elem('w:rPr');rp.append(elem('w:sz',val=22));rp.append(elem('w:szCs',val=22))
   if ri==0:rp.append(elem('w:b'))
   rr.append(rp);
   for li,line in enumerate(str(val).split('\n')):
    if li:rr.append(elem('w:br'))
    tt=elem('w:t');tt.text=line;rr.append(tt)
   pp.append(rr);tc.append(pp);tr.append(tc)
  e.append(tr)
 new.append(e);plain.append('| '+' | '.join(headers)+' |');plain.extend('| '+' | '.join(map(str,row))+' |' for row in rows)

p('问题四模型建立与求解','h1')
p('为解释开源模型能力随时间提升的来源，本文建立参数量与训练数据量参与的结构中介模型，将历史能力变化分解为规模扩张和非规模进步两部分。在此基础上，将问题三的资源最优配置接入能力分位响应，预测算力增速放缓情景下的未来前沿；损失—能力直接桥接与大样本榜单趋势作为独立对照。')
p('研究对象与能力指标','h2')
p('历史归因以训练元数据完整的基础模型为研究对象。附件C1经许可证、日期、参数量和六项任务完整性筛选，并按仓库去重后保留2449个模型；其中与C4确定性匹配且具有完整参数量、训练量和发布日期的基础模型为35个，来自10家机构。该子样本用于历史归因和结构前沿，2449条榜单记录用于分类型趋势对照。')
p('“开源”按附件中开放权重且许可证明确允许至少研究使用的口径筛选。主样本另要求C4开放权重标记为Yes、未标注底模且非明显混合专家模型。名称匹配保留参数量小数点，并核对数量级及发布机构；有歧义的记录不进入主样本。基础、对话／微调、合并和继续预训练模型分别处理，缺失的训练数据量不予填补。')
p('以IFEval、BBH、MATH Lvl 5、GPQA、MuSR和MMLU-PRO六项归一化得分的等权均值表示模型能力。记第i个模型在第k项任务上的得分为sᵢₖ，则')
eq(880)
p('为使能力响应能够在实数域内估计，将能力分转为logit尺度；仅在变换时将端点比例截断至[0.001,0.999]，原始得分保持不变：')
eq(883)
p('历史归因的时间T按发布日期距2023年1月1日的天数除以365.25计算，参数量N和训练量D均以十亿为单位，公式中的下标B表示这一单位。榜单趋势采用提交日期并以2024年6月1日为原点，算力季度趋势以2022年1月1日为原点。三类时间变量分别估计，不混用截距。C3的4573条排行榜来源记录与26条历史文献记录分层核对，不将不可比的早期得分接入能力趋势。')
p('C8用于核查逐任务得分。每个目录选择时间最新且可解析的评测记录，MuSR的谋杀推理、物体放置、团队分配三项任务分别以1/2、1/5、1/3为随机猜测基线。对选项长度归一化准确率a，先按100max{0,(a−b)/(1−b)}扣除基线，再将三项结果等权平均；仅三项齐全时计算聚合值。扫描1958个JSON发现4个损坏文件，其中1个目录存在有效替代，3个目录退出逐任务分析。匹配C1的1012个模型中，967个重建MuSR得分与榜单差异不超过10⁻⁶分。重建结果用于版本核对，不覆盖C1。')
p('规模与非规模进步的结构中介模型','h2')
p('模型路径与基本假设','h3')
p('时代变化一方面通过增加参数量和训练数据量改善能力，另一方面通过架构、优化和数据工程等因素改变相同规模下的能力。由此设置T→(lnN,lnD)→Z→S和T→Z→S两条路径：前者称为规模路径，后者在规模与评测口径可比的假设下解释为非规模进步。历史数据缺少模型级质量Q和配比p，故不将二者的历史贡献单独识别。')
# Keep the editable chapter fully self-contained; the existing publication diagram is inserted below.
FIG_MARK=len(new)
p('两维规模残差允许相关，以反映参数与数据投资的共同来源；主模型假定规模扰动与能力扰动独立，且残差分布跨期稳定。路径的因果解释还依赖未处理混杂和选择偏差足够小、比较时期规模范围具有重叠。上述条件作为解释假设，并在求解后通过结构对照与混杂压力测试检验其影响。')
p('规模中介与能力响应','h3')
p('令Mᵢ=(lnNᵢ,lnDᵢ)ᵀ，首先建立时间对两类规模投资的响应：')
eq(926)
p('其中a₀与aT为二维系数向量，E(εM,ᵢ|Tᵢ)=0，Cov(εM,ᵢ)=ΣM。沿用问题二已估计的A、B、α、ν，将参数量与训练量对可约损失的作用合成为规模指数：')
eq(932)
p('r表示总损失中可由规模扩张降低的部分，x越大表示对应的可约损失越低。固定问题二的标度形状，可减少35个主样本中待估计参数的数量，但需要接受该形状跨模型族迁移的假设。能力响应方程为')
eq(935)
p('β₀、κ和γT由附件C重新估计。κ刻画规模指数与能力的关系，γT刻画控制规模后的时间变化；历史均值模型不对两项系数施加正号约束。规模通过可约损失非线性作用于能力，且logit转回能力分也为非线性过程，因此有限时期贡献不能直接由回归系数的比例确定。')
p('参数估计与反事实分解','h3')
p('中介方程和能力方程均采用无权最小二乘法，分别最小化两维对数规模残差平方和及能力logit残差平方和。拟合后保留每个模型的成对规模残差及能力残差。普通能力预测先将每个训练期能力残差加到线性预测量上，经式（8-2）的expit逆变换回到分数尺度，再对所得分数取平均，以保留非线性变换下的噪声影响。')
p('为衡量初期t₀至后期t₁的变化，分别切换规模和非规模路径。定义G(t,t′)：非规模时间取t、规模中介时间取t′时的平均能力。将第a组规模残差代入中介方程，得到相应的反事实规模：')
eq(961)
p('上式中的指数逐分量作用。再将nM组成对规模残差与nY个能力残差交叉平均：')
eq(963)
p('两类残差来自同一拟合样本，故本次nM=nY=n；交叉平均实现跨方程独立的主假设，并保留N与D的残差相关性。四个角点G(t₀,t₀)、G(t₀,t₁)、G(t₁,t₀)、G(t₁,t₁)分别对应初期基准、仅改变规模、仅改变非规模条件、两条路径均转为后期。')
p('非线性响应下，切换两条路径的先后顺序会影响单项增量。因此对两种顺序取平均，定义规模贡献与非规模贡献为')
eq(969);eq(971)
p('两式相加恰为G(t₁,t₁)−G(t₀,t₀)，从而在能力分尺度完成可加分解。以各项贡献除以总变化得到贡献占比；总变化绝对值不超过10⁻⁶分时不计算比例。比较日期取主样本发布日期的上下四分位，机构重抽样时保持该区间固定。按发布机构整组有放回抽样400次，每次重估两组方程并重算贡献，以2.5%和97.5%分位给出95%区间。')
p('历史归因结果与稳健性检验','h2')
p('模型求解与贡献估计','h3')
p('35个基础模型的发布日期覆盖2021年3月21日至2024年9月24日。最小二乘估计得到lnN的截距和时间系数分别为1.322799、0.286528，lnD对应系数为6.410831、1.239217；能力方程的截距、规模系数及时间系数分别为−3.541544、1.217559、0.191554。按既定时间原点，拟合关系为')
estimate_equation()
p('在2023年4月3日至2024年6月7日的四分位比较区间内，反事实平均能力由9.115分增至14.273分，总变化为5.159分。表8-1给出两条路径的贡献；这些数值是同一扰动分布下的模型反事实变化，而非两个年度榜单均分直接相减。')
table('历史能力变化的路径分解',['路径','贡献／分','95%区间／分','占比'],[
 ['规模扩张','2.939','[0.359, 7.412]','56.98%'],['非规模进步','2.219','[−1.754, 5.887]','43.02%'],['总变化','5.159','[1.905, 10.340]','100%']], [2100,1700,3500,1772])
p('点估计中两条路径均为正，规模贡献略高。但非规模贡献区间包含零，其占比区间为−23.91%至89.50%；现有样本不能精确确定非规模贡献的符号和比例。机构重抽样反映机构内相关造成的估计波动，并不消除机构资源与能力之间的混杂。')
p('样本外验证与敏感性','h3')
p('采用整机构留出和发布日期阻塞留出，分别考察跨机构泛化及较早模型对较晚模型的预测。各训练折重新估计能力方程和残差分布，以分数尺度的平均绝对误差MAE及均方根误差RMSE评价；主模型与自由lnN、lnD系数模型、仅规模模型、仅时间模型使用相同划分。')
table('历史能力模型的样本外误差',['模型','机构留出\nMAE／RMSE','日期留出\nMAE／RMSE'],[
 ['标度响应＋时间（主模型）','3.205／4.646','6.426／8.810'],['自由lnN、lnD＋时间','3.776／5.139','6.657／9.136'],['仅规模','3.375／4.817','7.114／9.845'],['仅时间','5.506／7.158','9.088／12.075']],[3600,2736,2736])
p('主模型在这两种划分下的误差均低于三个对照，但日期留出误差较大，平均偏差为−4.387分，表明对后期能力存在低估。机构留出为35条预测；日期留出为54条预测，同一模型可在多个截点出现，不能将54条记录视为54次独立实验。')
p('进一步以自由规模模型检验标度结构的影响，其规模贡献占比降为44.44%，说明归因对响应形式敏感。对未观测混杂，先拟合x=ax+bxT+v，设v与能力扰动的相关参数为ρ，并按下式进行系数偏差压力测试：')
eq(996)
p('式中sx为缩减规模方程残差标准差，sY为能力方程残差标准差；截距同步调整为β₀(ρ)=β̂₀+[κ̂−κ(ρ)]ax，能力经验残差分布保持不变。取ρ∈[−0.6,0.6]时，规模份额为20.12%—93.09%。该范围表示假设混杂强度下的敏感性，不是份额置信区间。另将D整体或仅对较晚模型乘以0.5、1、2，并删除单项任务重估，以区分训练量测量误差和能力指标选择的影响。')
p('损失与能力得分的分层桥接','h2')
p('结构能力方程利用规模形状解释历史能力，直接桥接则检验损失数值能否换算成能力得分。两者证据不同。采用C6中的Val_Loss与LB_Average建立配对关系，LB_Average与其六任务均分在数值上相符；按Loss_Comparability分为高可比和中可比两层。C5用于重叠记录核对，中可比记录保留报告来源。')
p('对层ℓ的样本集合Iℓ，拟合损失增加时能力不增加的等渗回归，即在单调约束下最小化平方误差：')
eq(1010)
p('所得分段单调函数仅在该层观测损失的最小值与最大值之间提供预测，超出范围记为缺失。逐条留出时，只统计测试损失仍在训练折范围内的误差，并报告有效点数；中可比层还按报告来源整组留出，以减少同源记录造成的乐观偏差。')
table('直接桥接的有效验证范围及误差',['可比层','总点数','有效验证点','MAE／RMSE'],[
 ['高可比','7','5','0.370／0.444'],['中可比','68','67','7.276／9.039']],[2300,1300,2000,3472])
p('高可比层在局部范围内误差较小，但问题三未来配置的损失落在1.843—1.870，全部超出其支持范围。中可比层按报告整组留出时，66个有效预测的MAE为7.365分。因此高可比桥接不能直接支持本问前沿预测，中可比桥接仅作为跨报告情景对照，不能替代结构模型的迁移假设。')
p('算力增速放缓下的能力前沿预测','h2')
p('未来算力路径与资源配置','h3')
p('预测起点取C1最后有效提交日2025年3月13日。C4算力样本限定为2022年1月1日至预测起点之间发布、权重开放、训练算力为正、未标注底模、非明显混合专家且置信标记非Speculative的语言建模、代码生成或对话模型，共215个。按季度计算训练算力90%分位，采用普通最小二乘拟合其对数时间趋势；以最近一年算力90%分位作为起点预算。')
p('估计起点预算C⋆=3.303422×10²⁴ FLOPs，历史年对数增速gC=1.033713。令h为预测月数、rC为增速保留比例，则')
eq(1026)
p('rC=0、1/4、1/2分别对应算力冻结、历史增速四分之一及二分之一；主情景取1/4。对每个预算调用问题三求解器，固定参考配比p₀、域内最高50%文档质量切片、sQ=1、4096上下文及指数质量成本，得到(N*h,D*h,u*h,L*h)。损失函数和训练、注意力、提质三项成本均沿用问题三；其余质量上限与成本用于独立敏感性比较。')
p('条件前沿响应与技术趋势','h3')
p('历史归因描述平均能力变化，前沿预测改用完整N、D样本的条件90%分位。给定规模指数x和时间T，建立')
eq(1038)
p('κτ≥0保证固定时间条件下能力前沿不随规模指数增大而下降；截距和时间系数不受符号限制。将分位残差拆为非负的eᵢ⁺和eᵢ⁻，采用线性规划求解：')
eq(1040)
p('完整样本估计得到b₀,τ=−2.978248、κτ=1.123154、γT,τ=0.292187。分位响应与历史均值响应分别估计，其系数不作混用。')
p('为限制短期时间趋势的长期延伸，设技术推进速度按半衰期H衰减。令δ=12ln2/H，则h个月累计有效推进年数为')
eq(1044)
p('主情景取H=24个月，另以12个月和不衰减作敏感性。H为外推假设，不由现有短时间榜单识别。起点最优配置对应的能力logit为')
eq(1048)
p('其中t⋆按2023年1月1日为原点换算。主资源路径上质量上限均有效且配比固定，共同质量和配比乘子在可约损失比中抵消，故未来前沿为')
eq(1052)
p('该式将可约损失的相对下降转换为规模收益，再加入衰减后的非规模时间收益；需要L*h>E。它对应资源最优路径上的基础模型条件90%分位，不是最高分保证。质量与配比已进入资源模型，不在能力截距之外重复加分。若改变质量或配比导致共同乘子随时间变化，则相对损失收益也包含这些变化，不能全部解释为纯规模进步。')
p('同时按机构重抽样完整能力样本和算力元数据400次，重估分位系数、预算起点及算力趋势，构造条件预测的95%区间。问题二参数在该过程中固定。预算重抽样采用固定质量上限的快速求解，并在5.88×10²³—6.05×10²⁵ FLOPs的15个预算点与完整求解器比较：质量状态均为R2，损失相对误差为0。主资源路径另通过预算守恒检验，最大相对残差为2.51×10⁻¹⁶。')
p('前沿预测结果及支持范围','h3')
table('主算力放缓情景下的资源配置与能力前沿',['目标日期','预算／10²⁴\nFLOPs','N／十亿','D／十亿\nToken','损失','前沿／分\n及95%区间'],[
 ['2026-03-13','4.278','75.102','8350.038','1.863','45.83\n[22.89, 74.76]'],['2027-03-13','5.539','84.394','9622.229','1.856','51.30\n[23.06, 84.97]']],[1600,1400,1200,1500,1000,2372])
p('主情景中可约损失随资源增长下降，12个月与24个月的条件前沿分别为45.83和51.30分。两期区间较宽，反映小样本分位响应与预算趋势的不确定性。主样本最新发布日期早于预测起点约0.465年，未来配置还存在规模外推，因此该结果应解释为从2025年3月13日起算的条件情景，不能作为从论文完成日起算的实时预测。')
p('条件90%分位模型的机构留出覆盖率为82.86%，发布日期留出覆盖率为70.37%，对应logit分位损失为0.117和0.149，均显示样本外标定不足。名义90%条件分位与表中95%重抽样区间是不同概念，前者不是已达到90%覆盖的保证。')
p('对表8-4的两期损失，直接桥接的高可比层均返回缺失；中可比层两期预测均为27.51分，区间分别为[20.21,34.99]和[20.56,34.35]。等渗函数在该损失范围出现平台，因此不会像结构前沿一样随时间增长。两条路径的差异反映响应形式、来源可比性及时间项假设，不对二者作简单平均，也不向直接桥接额外叠加时间收益。')
p('大样本分类型对照与短期回测','h2')
p('为利用训练量不完整的榜单记录，对基础、对话／微调、合并模型分别建立仅含参数量和提交时间的90%分位对照：')
eq(1064)
p('其中g表示类型，bN,g≥0；时间以2024年6月1日为零点。由于缺少D，bT,g仅代表类型内的描述性时间变化，不用于非规模贡献归因。以最近三个月的对数参数分布为参照，无近期样本时使用全部训练记录；未来对数参数分布平移量取ν/(α+ν)·ln(Ch/C⋆)。该比例是此对照模型的外推规则，不是含提质成本时问题三的一般精确解。')
p('分别在参照参数分布和拟合残差分布上取101个中点分位，概率为(j−1/2)/101，j=1,…,101。将两组分位点交叉组合，按24个月半衰期推进时间，先得到能力预测分布，再取其90%分位；这样保留规模分布和残差分布对总体前沿的共同影响。')
p('滚动回测按历史月设定截点，仅使用此前提交的同类型模型及此前发布的算力记录，预测其后1—3个月的月度能力90%分位。目标月至少有5个模型才计入误差，朴素基线为上月同类型90%分位保持不变。')
table('分类型前沿的短期回测与长期情景',['模型类型','回测MAE\n模型／上月延续','12个月前沿\n及95%区间','24个月前沿\n及95%区间'],[
 ['基础模型','9.443／7.498','22.72\n[12.16, 47.59]','26.46\n[11.70, 63.21]'],['对话／微调','1.419／1.990','55.69\n[49.12, 60.68]','66.81\n[57.31, 72.75]'],['合并模型','4.041／3.457','62.49\n[53.34, 69.48]','75.02\n[63.32, 82.80]']],[1700,2400,2486,2486])
p('每类回测均有11条有效预测记录，只有对话／微调模型优于上月延续，MAE下降约28.7%。基础和合并类型未胜出，其长期情景仅作对照。表8-5给出同类型模型分布的总体90%分位，而表8-4给出最优资源路径上的条件90%分位，二者统计对象不同；基础模型的两种预测不能据数值差异直接判为矛盾。短期回测也不能替代12／24个月预测的长期实证检验。')
p('模型评价','h2')
p('本文通过规模中介方程和标度约束的能力响应，使历史贡献与前三问的资源关系保持一致，并将平均归因、条件前沿和总体分位预测分别估计与检验。数值核验通过贡献可加性、同一时间零效应、反向比较符号、关闭路径零贡献、规模导数有限差分、预算守恒、桥接出界处理及逐任务聚合等24项检查，说明实现与公式一致。')
p('解释范围仍由数据和假设决定。历史归因仅适用于35个元数据完整的基础模型，非规模份额对响应结构及未观测混杂敏感；机构重抽样不能消除这种偏差；元数据为事后快照，发布日期留出属于回溯队列检验。前沿预测依赖标度形状跨族迁移、技术半衰期和问题三质量上限情景，95%区间不包含所有上游结构误差，也不能保证实际Token供给。高可比桥接的支持缺口、条件前沿覆盖不足及部分类型回测失败均保留在结论中。因此，本问提供的是可核验的条件归因与情景前沿，而非适用于全部开源模型的固定技术贡献比例。')

# Add the existing publication figure without rewriting any original media or relationships.
PNS='http://schemas.openxmlformats.org/package/2006/relationships'
rels=E.fromstring(parts['word/_rels/document.xml.rels'])
ids={x.get('Id') for x in rels};rid='rIdQ4Rewrite'
assert rid not in ids
rel=E.SubElement(rels,'{'+PNS+'}Relationship');rel.set('Id',rid);rel.set('Type',NS['r']+'/image');rel.set('Target','media/q4_rewrite_structure.png')
parts['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['word/media/q4_rewrite_structure.png']=(ROOT/'figures/q4_causal_publication.png').read_bytes()
# Use python-docx only as a local drawing XML constructor, not to resave the reference package.
from docx import Document
from docx.shared import Inches
picdoc=Document();pp=picdoc.add_paragraph();pp.add_run().add_picture(str(ROOT/'figures/q4_causal_publication.png'),width=Inches(6.3))
pe=E.fromstring(E.tostring(pp._p))
for x in pe.xpath('.//*[local-name()="blip"]'):x.set('{'+NS['r']+'}embed',rid)
for x in pe.xpath('.//*[local-name()="docPr"]'):x.set('id','9001');x.set('name','问题四结构中介路径');x.set('descr','时间通过参数量与训练数据量的规模路径、以及非规模路径影响能力。')
pr=pe.find('w:pPr',NS)
if pr is None:pr=elem('w:pPr');pe.insert(0,pr)
pr.append(elem('w:pStyle',val='-'));pr.append(elem('w:jc',val='center'));pr.append(elem('w:keepNext'))
# caption shares original centered body-derived caption geometry, without heading numbering.
p('图8-1  规模路径与非规模路径的结构关系','caption');cap=new.pop();plain.pop()
new[FIG_MARK:FIG_MARK]=[pe,cap]
# Make retained first chapters exactly unchanged and preserve terminal section settings.
for x in old[866:-1]:body.remove(x)
for e in new:body.insert(len(body)-1,scrub(e))
parts['word/document.xml']=E.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
settings=E.fromstring(parts['word/settings.xml']);uf=settings.find('w:updateFields',NS)
if uf is None:uf=elem('w:updateFields');settings.append(uf)
uf.set(tag('w:val'),'true');parts['word/settings.xml']=E.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=E.fromstring(parts['[Content_Types].xml']);CNS='http://schemas.openxmlformats.org/package/2006/content-types'
if not any(x.get('Extension')=='png' for x in ct):
 d=E.SubElement(ct,'{'+CNS+'}Default');d.set('Extension','png');d.set('ContentType','image/png');parts['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
 for i in infos:z.writestr(i,parts.pop(i.filename))
 for name,data in parts.items():z.writestr(name,data)
# Structural scope and evidence gates.
with ZipFile(SRC) as zs,ZipFile(OUT) as zo:
 a=E.fromstring(zs.read('word/document.xml')).find('w:body',NS);b=E.fromstring(zo.read('word/document.xml')).find('w:body',NS)
 assert all(E.tostring(x)==E.tostring(y) for x,y in zip(list(a)[:866],list(b)[:866]))
 assert E.tostring(a[-1])==E.tostring(b[-1])
 changed=[name for name in zs.namelist() if zs.read(name)!=zo.read(name)]
 assert set(changed)<= {'word/document.xml','word/settings.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
 text=''.join(b.xpath('.//w:t/text()|.//m:t/text()',namespaces=NS))
 for phrase in ['尚需从外部模块核准','材料所述笛卡尔积','本版程序选择','机构混杂已经消除']:
  assert phrase not in ''.join(''.join(e.xpath('.//w:t/text()',namespaces=NS)) for e in list(b)[866:])
 assert zo.testzip() is None
assert hashlib.sha256(SRC.read_bytes()).hexdigest()==old_hash
notes=ROOT/'Q4/论文第四问重写正文.md';notes.write_text('\n\n'.join(plain),encoding='utf-8')
report={'source_sha256':old_hash,'output':OUT.name,'output_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'preserved_pre_q4_body_elements':866,'changed_parts':changed,'equations':eq_no,'tables':table_no,'source_unchanged':True,'model_recomputed':False,'chapter_text_characters':len(''.join(plain))}
(ROOT/'Q4/论文第四问重写核验.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))

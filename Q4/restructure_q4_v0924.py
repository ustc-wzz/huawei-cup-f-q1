"""Rebuild only Q4 from the user's supplied manuscript; numerical models are frozen.
Run with the bundled Python runtime. Outputs a standalone chapter and a full-paper copy.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
import csv, hashlib, json, re, sys
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/'Q4/paper_restructure_v0924';WORK.mkdir(exist_ok=True)
SRC=Path(sys.argv[1]) if len(sys.argv)>1 else Path('/Users/luochen/华为杯/F题/初稿/F题论文_问题四.docx')
OUT=ROOT/'F题论文_问题四主线调整_独立章节_v0.9.24.docx'
FULL=ROOT/'F题论文_问题四主线调整_完整副本_v0.9.24.docx'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
RNS='http://schemas.openxmlformats.org/package/2006/relationships'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def txt(e):return ''.join(e.xpath('.//w:t/text() | .//m:t/text()',namespaces=NS))
with ZipFile(SRC) as z:srcparts={n:z.read(n) for n in z.namelist()}
sroot=E.fromstring(srcparts['word/document.xml']);old=list(sroot.find('w:body',NS));assert txt(old[866])=='问题四模型建立与求解'
source_hash=sha(SRC)
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.5);sec.bottom_margin=Cm(2.5);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
# Unique styles prevent changes to any existing chapter when merged.
for name,size,bold in [('Q4Body',12,False),('Q4Title',16,True),('Q4H2',13,True),('Q4H3',12,True),('Q4Caption',10.5,False),('Q4Table',10.5,False),('Q4Equation',11,False)]:
 st=d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans SC');st.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
 pf=st.paragraph_format;pf.space_after=Pt(4);pf.line_spacing=Pt(18);pf.widow_control=True
 if name=='Q4Body':pf.first_line_indent=Pt(24)
 if name in ['Q4H2','Q4H3','Q4Title']:
  pf.keep_with_next=True;pf.space_before=Pt(10);pf.space_after=Pt(6)
  st._element.get_or_add_pPr().append(OxmlElement('w:outlineLvl'));st._element.pPr[-1].set(qn('w:val'),'0' if name=='Q4Title' else '1' if name=='Q4H2' else '2')
 if name in ['Q4Caption','Q4Equation','Q4Table']:pf.first_line_indent=Pt(0);pf.line_spacing=Pt(14.5)
# Explicitly suppress inherited title borders and the original template's line grid.
for name in ['Q4Body','Q4Title','Q4H2','Q4H3','Q4Caption','Q4Table','Q4Equation']:
 pr=d.styles[name]._element.get_or_add_pPr()
 snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pr.append(snap)
 if name=='Q4Title':
  bd=OxmlElement('w:pBdr')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');bd.append(e)
  pr.append(bd)
# A title style is used for the standalone chapter heading.
d.styles['Q4Title'].base_style=d.styles['Title']
d.styles['Q4Equation'].paragraph_format.line_spacing=1.0
content=[];eqmap={};used=[]
def add_inline(p,s):
 pattern=r'(?<![A-Za-zα-ωΑ-Ω])([A-Za-zα-ωΑ-Ω])_([A-Za-z0-9ℓα-ωΑ-Ω]+(?:,[A-Za-z0-9α-ωΑ-Ω]+)?)(⋆?)'
 last=0
 for m in re.finditer(pattern,s):
  if m.start()>last:p.add_run(s[last:m.start()])
  om=OxmlElement('m:oMath');sub=OxmlElement('m:sSubSup' if m.group(3) else 'm:sSub')
  for nm,val in [('e',m.group(1)),('sub',m.group(2))]+([('sup',m.group(3))] if m.group(3) else []):
   a=OxmlElement('m:'+nm);r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=val;r.append(t);a.append(r);sub.append(a)
  om.append(sub);p._p.append(om);last=m.end()
 if last<len(s):p.add_run(s[last:])

def p(s,style='Q4Body'):
 a=d.add_paragraph(style=style);add_inline(a,s);content.append({'type':'p','style':style,'text':s});return a
def h(s):
 a=p(s,'Q4H2')
 # Let headings and their following paragraphs flow naturally.
 return a
def h3(s):return p(s,'Q4H3')
def note(s):return p(s,'Q4Caption')
def formula(idx,num):
 """Preserve native Office Math nodes, restyle only their container and number."""
 if d.paragraphs:d.paragraphs[-1].paragraph_format.keep_with_next=True
 e=deepcopy(old[idx]);assert e.tag==qn('w:tbl')
 for b in e.xpath('.//w:bookmarkStart|.//w:bookmarkEnd|.//w:lastRenderedPageBreak',namespaces=NS):b.getparent().remove(b)
 # Tables contain 1 or 2 equation rows, each with its own number.
 nums=num if isinstance(num,list) else [num]
 for row,label in zip(e.findall('w:tr',NS),nums):
  cells=row.findall('w:tc',NS)
  for c in cells:
   for pp in c.findall('w:p',NS):
    pr=pp.find('w:pPr',NS)
    if pr is not None:pp.remove(pr)
    pr=E.Element(qn('w:pPr'));sty=E.SubElement(pr,qn('w:pStyle'));sty.set(qn('w:val'),'Q4Equation');pp.insert(0,pr)
  last=cells[-1]
  for c in list(last):
   if c.tag!=qn('w:tcPr'):last.remove(c)
  pp=E.SubElement(last,qn('w:p'));pr=E.SubElement(pp,qn('w:pPr'));st=E.SubElement(pr,qn('w:pStyle'));st.set(qn('w:val'),'Q4Equation');jc=E.SubElement(pr,qn('w:jc'));jc.set(qn('w:val'),'right');rr=E.SubElement(pp,qn('w:r'));E.SubElement(rr,qn('w:t')).text=f'({label})'
  trp=row.find('w:trPr',NS)
  if trp is None:trp=E.SubElement(row,qn('w:trPr'))
  E.SubElement(trp,qn('w:cantSplit'))
 # Set consistent full text width. Original mathematical layout remains editable.
 pr=e.find('w:tblPr',NS)
 if pr is not None:e.remove(pr)
 pr=E.Element(qn('w:tblPr'));ww=E.SubElement(pr,qn('w:tblW'));ww.set(qn('w:w'),'9072');ww.set(qn('w:type'),'dxa');e.insert(0,pr)
 borders=E.SubElement(pr,qn('w:tblBorders'))
 for k in ['top','left','bottom','right','insideH','insideV']:el=E.SubElement(borders,qn('w:'+k));el.set(qn('w:val'),'nil')
 for row in e.findall('w:tr',NS):
  cells=row.findall('w:tc',NS)
  for cell,width in zip(cells,[120,7882,1070]):
   cp=cell.find('w:tcPr',NS)
   if cp is None:cp=E.SubElement(cell,qn('w:tcPr'))
   cw=cp.find('w:tcW',NS)
   if cw is None:cw=E.SubElement(cp,qn('w:tcW'))
   cw.set(qn('w:w'),str(width));cw.set(qn('w:type'),'dxa')
 d.element.body.insert(len(d.element.body)-1,e)
 content.append({'type':'formula','numbers':nums,'source_body_index':idx,'text':txt(e)})
 eqmap[str(idx)]=nums

def table(caption,headers,rows,widths=None):
 cap=p(caption,'Q4Caption');cap.paragraph_format.keep_with_next=True
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Cm(w)
 for j,s in enumerate(headers):t.rows[0].cells[j].text=str(s)
 for row in rows:
  for c,s in zip(t.add_row().cells,row):c.text=str(s)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[j])
   for pp in c.paragraphs:
    pp.style=d.styles['Q4Table'];pp.paragraph_format.keep_with_next=i<len(t.rows)-1;pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(4)
    pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
    if i==0:
     for rr in pp.runs:rr.bold=True
   cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
   for k in ['top','bottom','left','right']:
    el=OxmlElement('w:'+k);el.set(qn('w:val'),'single' if (i==0 and k in ['top','bottom']) or (i==len(t.rows)-1 and k=='bottom') else 'nil');el.set(qn('w:sz'),'8' if i==0 and k=='top' else '4');bd.append(el)
   cp.append(bd)
 content.append({'type':'table','caption':caption,'headers':headers,'rows':rows})
 return t

def fig(path,caption,width=16):
 pp=d.add_paragraph(style='Q4Caption');pp.paragraph_format.line_spacing=1.0;pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.keep_with_next=True;pp.add_run().add_picture(str(ROOT/path),width=Cm(width))
 cp=p(caption,'Q4Caption');cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 used.append(path);content.append({'type':'figure','path':path,'caption':caption})

# Two additional equations are algebraic expansions of the existing model.
import subprocess
NEW_EQ={
 'loss_factor':r'L_h^\star-E=\left[A(N_h^\star)^{-\alpha}+B(D_h^\star)^{-\nu}\right]m_h,\qquad m_h=\exp(-\lambda u_h+\eta_p h(p_h)).',
 'scale_identity':r'x(N_h^\star,D_h^\star)-x(N_0^\star,D_0^\star)=-\ln\frac{L_h^\star-E}{L_0^\star-E}.'}
eqmd=WORK/'新增公式.md';eqmd.write_text('\n\n'.join('$$\n'+v+'\n$$' for v in NEW_EQ.values()))
eqdoc=WORK/'新增公式.docx'
subprocess.run(['/Users/luochen/miniconda3/bin/pandoc',str(eqmd),'-o',str(eqdoc)],check=True)
with ZipFile(eqdoc) as z:er=E.fromstring(z.read('word/document.xml'))
newmath=er.xpath('//m:oMath',namespaces=NS)
def newformula(key,num):
    if d.paragraphs:d.paragraphs[-1].paragraph_format.keep_with_next=True
    pp=d.add_paragraph(style='Q4Equation');pp.paragraph_format.tab_stops.add_tab_stop(Cm(15.8),__import__('docx').enum.text.WD_TAB_ALIGNMENT.RIGHT)
    pp._p.append(deepcopy(newmath[list(NEW_EQ).index(key)]));pp.add_run('\t('+num+')')
    content.append({'type':'newformula','numbers':[num],'latex':NEW_EQ[key]});eqmap[key]=[num]

forecast=list(csv.DictReader((ROOT/'output_q4_evolution/tables/coupled_frontier_forecast.csv').open()))
def frow(s,h):return next(r for r in forecast if r['scenario']==s and int(r['horizon_months'])==h and float(r['technical_half_life_months'])==24)
main_name='quarter_historical_growth'
freeze=[float(frow('compute_frozen',hh)['forecast']) for hh in [12,24]];main=[float(frow(main_name,hh)['forecast']) for hh in [12,24]]
assets='Q4/paper_restructure_v0924/assets/'
p('8  技术演进的历史归因与条件前沿预测','Q4Title')
p('本问围绕“规模扩张和非规模变化怎样共同影响能力，以及算力增速放缓后这种关系如何延伸”展开。历史部分估计时间经参数量与训练量传递的规模路径，以及控制规模后的非规模路径，再用反事实计算分解能力增量；预测部分将算力情景接入问题三的最优资源配置，计算规模收益并叠加非规模时间收益，得到条件能力前沿。两部分共享问题二的规模指数x(N,D)，历史均值响应与前沿分位响应的系数分别估计，历史贡献比例仅用于解释历史。')
h('8.1  研究目标、样本与能力指标')
p('历史归因和结构前沿使用参数量、训练Token数及发布日期齐全的35个基础模型，覆盖10家机构；2449条经筛选、去重的榜单记录用于分类型趋势对照。前者回答完整规模条件下的历史变化与资源前沿，后者检查更广泛模型类型的描述性趋势。“开源”按附件声明权重开放且许可证明确允许至少研究使用的口径筛选，主样本排除已标注底模的衍生模型及明显混合专家模型，不填补缺失训练量。')
p('综合能力S为IFEval、BBH、MATH Lvl 5、GPQA、MuSR和MMLU-PRO六项归一化得分的等权均值，单位为分：')
formula(872,'8-1')
p('将能力分转换为Z=logit(S/100)，其中logit为比例p的对数优势ln[p/(1−p)]；计算变换时将比例截断至[0.001,0.999]，逆变换为S=100/(1+e⁻ᶻ)。N、D分别表示以十亿计的参数量与训练Token数，公式下标B表示该单位。历史时间T为发布日期距2023年1月1日的年数，每年按365.25天计算；榜单对照采用提交时间。数据匹配、来源分层与逐任务得分已作核查，相关规则、异常文件及其他时间原点集中列于附录四A。')
h('8.2  历史能力变化的结构模型')
p('结构中介模型＝用中间规模变量刻画时间影响能力的传递路径。如图8-1，时间通过参数量与训练量变化影响能力，同时保留直接进入能力响应的时间路径。控制规模后的剩余时间变化，在评测口径可比、未处理混杂与选择偏差足够小且比较时期规模充分重叠的假设下，解释为非规模进步；架构、优化与数据工程等因素合并体现，现有数据不再分别识别历史质量和配比贡献。')
fig(assets+'q4_causal_publication.png','图8-1  历史能力变化的规模路径与非规模路径',15.5)
p('先用两维对数规模向量M描述时间对资源投入的影响：')
formula(885,'8-2')
p('a₀与a_T为待估计的二维系数，ε_M为规模扰动；允许同一模型的参数量与训练量扰动相关。随后承接问题二已估计的A、B、α、ν，构造可约损失r及规模指数x：')
formula(887,'8-3')
p('x越大，表示该规模对应的可约损失越低。固定标度形状可以减少小样本的待估参数，同时要求接受这一形状跨模型族迁移的假设。能力响应为')
formula(889,'8-4')
p('β₀、κ、γ_T均在本问的35个基础模型上重新估计，历史模型不限制κ与γ_T的符号。主模型假定两维规模扰动与能力扰动ε_Y相互独立，且扰动分布跨期稳定。规模响应和能力逆变换均具有非线性，有限时期的贡献须在能力分尺度计算，不能由两项回归系数的比例直接得到。')
h('8.3  历史贡献的求解、结果与验证')
h3('8.3.1  参数估计与四种反事实情境')
p('求解依次为：最小二乘估计方程 → 保留成对规模残差与能力残差 → 计算四种情境的平均能力 → 求两类贡献 → 按机构重抽样计算区间。两组方程均采用无权最小二乘，能力响应的截距、规模系数和时间系数分别为−3.541544、1.217559和0.191554；中介系数见附录四A。')
p('反事实计算＝在模型中固定一条路径的状态，只切换另一条路径，比较平均能力。令G(t,t′)表示非规模条件取t、规模条件取t′时的平均能力：在t′下生成参数量与训练量，在t下计算能力，再对经验扰动平均。N与D的残差按原模型成对保留，并与能力残差交叉平均，以实现跨方程独立的假设；完整双重求和见附录四A。')
table('四种反事实情境及其含义',['非规模状态／规模状态','初期规模 t₀','后期规模 t₁'],[['初期非规模状态 t₀','G(t₀,t₀)\n初期基准','G(t₀,t₁)\n只改变规模'],['后期非规模状态 t₁','G(t₁,t₀)\n只改变非规模条件','G(t₁,t₁)\n两条路径均改变']],[5.0,5.5,5.5])
p('分别按“先改变规模”和“先改变非规模条件”计算增量，再对两种顺序取平均。由此定义规模贡献Δ_S和非规模贡献Δ_T：')
formula(899,['8-5','8-6'])
p('两类贡献之和恰为G(t₁,t₁)−G(t₀,t₀)，使总能力变化完整分配而不重复计算。比较日期固定为主样本发布日期的上下四分位，即2023年4月3日与2024年6月7日；按机构整组有放回重抽样400次，每次重估两组方程及经验残差，使用2.5%和97.5%分位构造95%区间。')
h3('8.3.2  贡献结果与稳健性')
table('表8-1  2023年4月3日至2024年6月7日的历史能力变化分解',['路径','贡献／分','95%重抽样区间／分','占比'],[['规模扩张','2.939','[0.359, 7.412]','56.98%'],['非规模进步','2.219','[−1.754, 5.887]','43.02%'],['总变化','5.159','[1.905, 10.340]','100%']],[3.5,3,6,3.5])
p('反事实平均能力由9.115分升至14.273分。表8-1的非规模贡献区间包含零，43.02%是有较大不确定性的点估计，其份额95%区间为[−23.91%,89.50%]；机构重抽样反映机构内相关引起的估计波动，不能消除资源与能力之间的混杂。表内由未舍入值计算，显示值相加可能存在末位差异。')
fig(assets+'fig_39_q4_mediation_and_sensitivity.png','图8-2  历史贡献区间与假定混杂强度下的份额变化',16)
p('自由规模响应下的规模份额为44.44%；当假定混杂相关系数ρ在[−0.6,0.6]内变化时，规模份额为20.12%—93.09%（图8-2）。后者是压力测试范围，不是置信区间，说明“规模是否占多数”依赖响应结构与混杂假设。训练量测量误差的敏感性设置见附录四B。')
table('表8-2  历史能力模型在相同划分下的样本外误差',['模型','机构留出\nMAE／RMSE','日期留出\nMAE／RMSE'],[['标度响应＋时间','3.205／4.646','6.426／8.810'],['自由lnN、lnD＋时间','3.776／5.139','6.657／9.136'],['仅规模','3.375／4.817','7.114／9.845'],['仅时间','5.506／7.158','9.088／12.075']],[6.4,4.8,4.8])
p('MAE＝平均绝对误差，RMSE＝均方根误差，均以能力分计。表8-2与图8-3显示主模型在两种留出方式下均优于三个对照，支持同时保留规模与时间项；但机构留出MAE相对仅规模只降低0.170分，日期留出平均偏差为−4.387分，向较晚模型迁移时仍明显低估。这些预测检验不直接证明贡献比例的因果确定性。')
fig(assets+'historical_validation.png','图8-3  历史能力模型的机构留出与发布日期阻塞留出误差',16)
h('8.4  算力放缓下的前沿模型与求解')
h3('8.4.1  预算情景与最优资源配置')
p('历史模型解释平均能力变化；本节另行估计前沿分位响应，共享规模指数x的定义与问题二参数，系数分别拟合。条件第90百分位＝给定规模与时间时能力分布的第90百分位，用于表示领先水平。计算链为：算力增长情景 → 问题三最优资源配置 → 规模收益变化 → 加入非规模时间收益 → 条件能力前沿。')
p('预测以附件最后有效提交日2025年3月13日为基点，12个月和24个月分别对应2026年3月13日、2027年3月13日。215条开放权重算力记录的季度90%分位用于拟合对数趋势，最近一年算力90%分位给出C⋆=3.303422×10²⁴ FLOPs，历史年对数增速g_C=1.033713。令h为预测月数，预算情景为')
formula(935,'8-7')
p('r_C=0、1/4、1/2分别对应算力冻结、历史对数增速保留四分之一及二分之一，主情景取1/4。每个预算沿用问题三的损失函数和训练、注意力、提质成本，固定参考配比p₀、域内最高50%文档质量参照、s_Q=1、4096上下文及指数质量成本，求得最优参数量、训练量、质量增量和损失。上述资源情景沿用同一D成本代理，实际Token供给仍属于外推假设。')
h3('8.4.2  从最优损失到规模收益')
p('在35个完整N、D样本上建立条件分位响应：')
formula(939,'8-8')
p('分位回归＝通过非对称绝对残差损失估计指定分位水平。采用线性规划求解，κ_τ≥0保证固定时间下前沿随规模指数单调不减；估计的截距、规模系数与时间系数分别为−2.978248、1.123154和0.292187，线性规划的残差拆分见附录四C。')
p('问题三的最优损失减去不可约损失E后，等于规模项乘以质量和配比的共同修正项：')
newformula('loss_factor','8-9')
p('式中λ=θ_Q/s_Q，主情景设s_Q=1，η_p承接问题二；λ与质量尺度未由数据分别识别。u_h为质量增量，h(p_h)为沿用问题二的配比响应函数；h(p_h)与表示预测月数的下标h含义不同。主资源路径的配比始终固定，质量持续达到同一参照上限，且E、A、B、α、ν保持不变，因此m_h=m₀。将式（8-9）在未来与起点相除，再代入x=−ln r，得到关键连接：')
newformula('scale_identity','8-10')
p('式（8-10）把问题三输出的相对可约损失下降转换为问题四的规模指数增益；乘以κ_τ即可进入前沿的logit响应。它要求L_h⋆、L₀⋆均大于E。若质量或配比随时间改变而使共同乘子变化，该损失比还包含相应变化，便不能全部解释为规模收益。')
h3('8.4.3  非规模时间推进与最终前沿')
p('为限制短期非规模趋势的长期线性延伸，设推进速度按H个月的半衰期衰减。令δ=12ln2/H，h个月的有效推进年数为')
formula(944,'8-11')
p('主情景H=24个月，另考察12个月与不衰减；半衰期是情景设定。记z⋆为起点最优资源对应的分位响应logit，按起点规模与时间代入式（8-8）计算，则未来前沿由“起点能力＋规模收益＋非规模时间收益”组成：')
formula(948,'8-12')
p('expit(z)=1/(1+e⁻ᶻ)。式（8-12）的中间项由式（8-10）给出，末项加入衰减后的非规模时间收益；历史贡献比例不进入该式。按机构重抽样完整能力样本和算力元数据400次，重估分位系数、预算起点与增长趋势，构造95%重抽样区间，问题二参数固定。快速求解器和预算一致性核验见附录四C。')
fig(assets+'frontier.png','图8-4  以2025年3月13日为基点的三种算力情景条件前沿',12.5)
note('图注：技术半衰期固定为24个月；阴影为主情景95%重抽样区间。条件第90百分位为预测对象，95%为估计区间水平。起点亦为模型估计。')
h3('8.4.4  前沿结果与覆盖率检验')
rows=[]
for hh in [0,12,24]:
 r=frow(main_name,hh);rows.append([r['target_date']+('（起点）' if hh==0 else ''),f"{float(r['compute_budget'])/1e24:.3f}",f"{float(r['N_B']):.3f}",f"{float(r['D_B']):.3f}",f"{float(r['loss']):.3f}",f"{float(r['forecast']):.2f}\n[{float(r['lo']):.2f}, {float(r['hi']):.2f}]"])
table('表8-3  以2025年3月13日为基点的主情景资源配置与条件前沿',['目标日期','预算\n10²⁴ FLOPs','N／十亿','D／十亿\nToken','损失','条件第90百分位估计\n及95%重抽样区间／分'],rows,[2.7,1.9,1.7,2.1,1.3,6.3])
p('表8-3的起点38.73分也是模型估计；主情景在12个月和24个月分别为45.83分、51.30分，相对起点增加7.10分、12.57分。两期参数量均超过主样本最大70B，训练量仍低于样本最大18000B；预测起点已晚于完整样本最新发布日期约0.465年，因此结果同时包含时间与参数规模外推。')
rows=[]
for name,label in [('compute_frozen','算力冻结'),(main_name,'历史对数增速保留1/4'),('half_historical_growth','历史对数增速保留1/2')]:
 vals=[frow(name,hh) for hh in [12,24]];rows.append([label]+[f"{float(r['forecast']):.2f} [{float(r['lo']):.2f}, {float(r['hi']):.2f}]" for r in vals])
table('表8-4  固定24个月技术半衰期的算力放缓情景比较',['算力情景','12个月估计及95%区间／分','24个月估计及95%区间／分'],rows,[5.4,5.3,5.3])
p('24个月时，冻结、四分之一增速与二分之一增速分别给出49.07、51.30和53.52分。算力冻结时规模收益为零，继续增长来自模型中推进的非规模时间项；恢复部分算力增长后，主情景相对冻结增加2.23分，二分之一增速再比主情景增加2.22分。各情景区间较宽且重叠，点估计差异的确定性有限，不能据此断言实际前沿必然按该顺序分离。')
p('条件第90百分位的机构留出覆盖率为82.86%，发布日期留出为70.37%，均低于名义90%，对应logit分位损失为0.117和0.149，表明样本外标定不足。覆盖率＝实际能力不超过预测条件分位的比例，与表中95%重抽样区间的置信水平不同；该区间也未覆盖标度跨族迁移、上游参数和技术衰减设定的全部误差。')
h('8.5  直接桥接与分类型预测对照')
h3('8.5.1  另一条损失换算路线的支持范围')
p('主前沿通过相对损失变化与结构分位响应完成预测。本节另用C6配对的损失和六任务平均分，检查绝对损失直接换算为能力能支持到何种范围。按损失可比层ℓ分别拟合等渗回归＝在损失增加时能力不增加的约束下，使平方误差最小的单调函数：')
formula(925,'8-13')
p('预测只在各层观测损失的最小值与最大值之间定义。逐条留出时也仅对处于训练范围内的测试点计算误差，故须同时报告有效点数。')
table('表8-5  直接损失桥接的验证误差及未来支持范围',['可比层','总点数／\n有效留出点','MAE／RMSE\n分','对主情景未来损失的支持'],[['高可比','7／5','0.370／0.444','1.863、1.856均超出范围'],['中可比','68／67','7.276／9.039','可计算，来源与验证集混合']],[2.3,3.2,3.7,6.8])
p('高可比层观测范围为[2.0933,2.5978]，不能支持主情景未来损失1.863和1.856；中可比桥接在两期均给出27.51分，95%区间分别为[20.21,34.99]和[20.56,34.35]，相同点估计来自等渗函数的平台段。中可比层按来源字段留出的66条有效预测MAE为7.365分，且来源别名尚未归并。该结果是独立的跨报告换算对照，与结构前沿在响应、评测来源和时间假设上不同，不再叠加时间收益，也不与主前沿求平均。')
h3('8.5.2  分类型总体前沿与短期回测')
p('在2449条榜单记录中，分别对192个基础、1542个对话／微调及667个合并模型建立仅含参数量与提交时间的90%分位对照；因缺少D，时间项仅表示描述性变化。基于近期参数分布与经验残差分布，按同一算力主情景和24个月技术半衰期，生成类型总体能力分布，再取第90百分位。它描述类型总体的领先水平，与最优资源路径上的条件前沿分别解释；模型和分布构造见附录四C。')
fig(assets+'broad_forecast.png','图8-5  分类型总体能力第90百分位的长期情景对照',12.5)
note('图注：算力保留历史对数增速1/4，技术半衰期24个月；误差线为按机构重抽样能力样本得到的95%区间，算力趋势固定。三类型横坐标略作错开以避免误差线重叠。')
p('图8-5对应的12个月／24个月点估计依次为：基础22.72／26.46分，对话／微调55.69／66.81分，合并62.49／75.02分。为检验短期预测作用，滚动预测未来1—3个月的同类型月度90%分位，目标月至少有5个模型，以沿用上月90%分位作为基线。')
table('表8-6  分类型月度前沿的短期回测',['模型类型','模型MAE／分','上月延续MAE／分','有效预测记录'],[['基础','9.443','7.498','11'],['对话／微调','1.419','1.990','11'],['合并','4.041','3.457','11']],[4.4,3.6,4.6,3.4])
fig(assets+'broad_validation.png','图8-6  分类型短期预测与上月延续基线的误差比较',10.5)
p('表8-6与图8-6中，只有对话／微调类型优于上月延续，MAE减少约28.7%；基础与合并类型均未胜出，其长期结果仅作情景对照。该回测检验类型总体前沿，不直接检验问题三资源路径上的条件前沿，1—3个月回测也不能替代12—24个月的长期验证。')
h('8.6  主要结论与适用范围')
p('历史主样本的规模与非规模贡献点估计为56.98%和43.02%，非规模贡献区间包含零，比例对响应结构和混杂假设敏感。承接问题三最优资源后，算力保留历史对数增速四分之一的条件前沿为45.83分和51.30分；24个月冻结及二分之一增速对照分别为49.07分和53.52分，冻结条件下的增长来自非规模时间项。两条主分析共同依赖规模响应的迁移假设，历史比例不直接用于未来预测。')
p('结论适用于35个元数据完整的基础模型及设定的资源路径。前沿区间较宽、样本外覆盖不足，高可比直接桥接不支持未来损失，分类型短期回测也只在对话／微调类型胜过基线。因此，本问提供带有明确假设和不确定范围的历史归因与前沿情景；数值一致性检查及数据审计不替代因果识别与未来实测验证，完整核验清单见附录四D。')
ap=p('附录四  数据核查与模型补充','Q4Title');ap.paragraph_format.page_break_before=True
h('四 A  数据口径与经验扰动计算')
p('C1经许可证、日期、参数量和六任务完整性筛选，按仓库去重后保留2449个模型。与C4按规范化名称、参数数量级及发布机构确定性匹配；名称保留参数小数点，有歧义或无法核验的重上传退出主样本。完整基础样本35个，发布日期覆盖2021年3月21日至2024年9月24日。C3的4573条榜单来源记录与26条历史文献记录分层核对，不将不可比历史得分拼接进趋势。')
p('C8逐任务核查扫描1958个JSON，4个损坏文件中1个目录有可解析替代，另3个目录退出逐任务分析，原件保留。MuSR三项叶子任务使用随机猜测基线1/2、1/5、1/3，先计算100max{0,(a−b)/(1−b)}，再等权平均；a为选项长度归一化准确率。匹配C1的1012个模型中，967个重建MuSR与榜单差异不超过10⁻⁶分；结果仅用于核查，不覆盖C1。')
p('能力变换的完整定义为：')
formula(874,'四-1')
p('中介方程与能力方程采用无权最小二乘。lnN的截距、时间系数为1.322799、0.286528；lnD为6.410831、1.239217。能力拟合为：')
formula(904,'四-2')
p('普通能力预测对训练期能力残差分别作expit逆变换再平均。反事实计算首先使用第a组成对规模残差生成t′时刻规模：')
formula(894,'四-3')
p('随后将规模残差与能力残差交叉组合：')
formula(896,'四-4')
p('本次n_M=n_Y=35。交叉平均对应跨方程独立假设，同时保留N与D残差的相关性。四个角点分别为初期基准、仅改变规模、仅改变非规模条件、两条路径同时改变。总变化绝对值不超过10⁻⁶分时不计算份额。机构重抽样固定比较日期，每次重估两组方程和经验残差。')
h('四 B  稳健性检验与桥接细节')
p('机构留出形成35条预测；日期阻塞留出形成54条预测记录，涉及21个不同模型，同一模型可出现在多个截点。训练折重新估计能力系数和扰动分布，表8-2各模型使用相同划分。元数据为事后快照，该检验属于回溯发布队列检验。')
p('混杂压力测试先拟合x=a_x+b_x T+v，令ρ表示v与能力扰动的假设相关程度，并使用：')
formula(919,'四-5')
p('s_x、s_Y分别为上述规模残差与能力残差标准差；截距同步增加“原规模系数与调整后规模系数之差乘以a_x”，能力经验残差保持不变。ρ∈[−0.6,0.6]对应的份额范围属于敏感性情景，不是置信区间。对D整体或仅较晚模型乘以0.5、1、2，检验历史贡献对训练量测量误差的敏感性；逐项删除能力任务的实验则用于检查大样本描述性预测对指标选择的敏感性。')
p('桥接采用C6的Val_Loss、LB_Average及Loss_Comparability字段，C5只作重叠记录核查。高可比7点属于Pythia同族局部范围；中可比68点含跨验证集及近似报告损失。来源留出按Loss_Source原字符串分组；同一报告的不同写法尚未归并，因此不能声称已排除同报告信息相关。')

h('四 C  前沿求解与独立分类型对照')
p('算力样本限定在2022年1月1日至2025年3月13日之间发布，权重开放、算力为正、未标注底模、非明显混合专家且置信标记非Speculative的语言建模、代码生成或对话模型，共215个。季度算力90%分位采用对数时间最小二乘拟合，时间原点为2022年1月1日。')
p('条件分位响应通过以下线性规划求解，τ=0.9：')
formula(941,'四-6')
p('分位残差由非负e⁺、e⁻表示。κ_τ≥0，截距和时间系数不受符号限制。式（8-12）中的起点logit为：')
formula(946,'四-7')
p('t⋆按历史响应的2023年1月1日原点计算。机构及算力元数据重抽样400次，重估分位系数、预算起点与增长趋势，问题二参数固定。质量固定在同一上限的快速求解曾在5.88×10²³—6.05×10²⁵ FLOPs的15个预算点与完整求解器核对，均达到上限且损失相对误差为0；主路径预算最大相对残差为1.26×10⁻¹⁶。这些检查验证计算一致性，不能验证真实训练效果。')
p('大样本对照按模型类型g分别估计：')
formula(959,'四-8')
p('其中b_N,g≥0，榜单时间原点为2024年6月1日。最近三个月的对数参数分布为参照，无近期记录则采用全部训练记录；未来平移量为ν/(α+ν)·ln(C_h/C⋆)，这是对照模型的外推规则。分别在参数分布和拟合残差分布取101个中点分位，交叉组合后取预测分布90%分位，技术时间采用24个月半衰期。表四-1的区间仅按发布机构重抽样能力样本，算力趋势固定；与表8-3同时重抽样算力元数据的口径不同。')
table('表四-1  分类型总体前沿的长期情景对照',['类型','12个月前沿及95%区间 / 分','24个月前沿及95%区间 / 分'],[['基础','22.72 [12.16,47.59]','26.46 [11.70,63.21]'],['对话／微调','55.69 [49.12,60.67]','66.81 [57.31,72.75]'],['合并','62.49 [53.34,69.48]','75.02 [63.32,82.80]']],[3.2,6.4,6.4])
p('表四-1表示类型总体分布的90%分位，表8-3表示问题三最优资源路径上的条件90%分位；统计对象不同。仅对话／微调对照在既有短期回测中胜过上月延续，其他类型的长期数值仅保留为情景对照。')
h('四 D  数值与数据一致性检查')
p('已有保存核验共24项通过，按下表集中列出。检查对象是代数恒等式、边界处理、来源一致性和数值实现，不表示因果解释或未来预测已经得到实证验证。')
table('表四-2  已保存的24项数值与数据核验',['序号','核验内容','保存状态'],[['1', '参数小数点名称区分', '通过'], ['2', '主样本模型唯一', '通过'], ['3', '主样本均为基础模型', '通过'], ['4', '训练量未填补', '通过'], ['5', '贡献分解可加', '通过'], ['6', '同一时间零效应', '通过'], ['7', '反向时间贡献变号', '通过'], ['8', '规模系数为零时间接贡献为零', '通过'], ['9', '时间系数为零时直接贡献为零', '通过'], ['10', '规模弹性与有限差分一致', '通过'], ['11', '问题三预算可行', '通过'], ['12', '质量上限快速求解一致', '通过'], ['13', '预测分数在定义范围内', '通过'], ['14', '区间上下界有序', '通过'], ['15', '冻结算力规模收益为零', '通过'], ['16', '未来日期从最后有效提交日起算', '通过'], ['17', '算力数据无未来发布日期', '通过'], ['18', '不支持的桥接返回缺失', '通过'], ['19', '桥接函数单调', '通过'], ['20', 'MuSR三任务等权聚合', '通过'], ['21', 'MuSR随机基线归一化', '通过'], ['22', '损坏文件完整入账', '通过'], ['23', '输入文件哈希一致', '通过'], ['24', 'C8原件哈希一致', '通过']],[1.5,12,2.5])
for row in d.tables[-1].rows[1:]:
 for cell in row.cells:
  for pp in cell.paragraphs:pp.paragraph_format.keep_with_next=False

# Independent chapter footer.
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');rr._r.addnext(field)
d.save(OUT)
# Merge new chapter into the original package. All pre-Q4 body nodes stay byte-equivalent under c14n.
with ZipFile(OUT) as z:newparts={n:z.read(n) for n in z.namelist()}
nroot=E.fromstring(newparts['word/document.xml']);nbody=nroot.find('w:body',NS);fresh=[deepcopy(x) for x in nbody if x.tag!=qn('w:sectPr')]
rels=E.fromstring(srcparts['word/_rels/document.xml.rels']);newrels=E.fromstring(newparts['word/_rels/document.xml.rels']);idmap={}
existing={x.get('Id') for x in rels}
for rr in newrels:
 if rr.get('Type','').endswith('/image'):
  rid='rIdQ4v0924_'+rr.get('Id');assert rid not in existing
  oldtarget=rr.get('Target');newtarget='media/q4v0924_'+Path(oldtarget).name
  srcparts['word/'+newtarget]=newparts['word/'+oldtarget]
  el=deepcopy(rr);el.set('Id',rid);el.set('Target',newtarget);rels.append(el);idmap[rr.get('Id')]=rid
for el in fresh:
 for child in el.iter():
  for k,v in list(child.attrib.items()):
   if k.startswith('{'+NS['r']+'}') and v in idmap:child.set(k,idmap[v])
# Begin replacement at a clean chapter boundary.
pr=fresh[0].find('w:pPr',NS);E.SubElement(pr,qn('w:pageBreakBefore'))
body=sroot.find('w:body',NS)
for el in list(body)[866:-1]:body.remove(el)
for el in fresh:body.insert(len(body)-1,el)
styles=E.fromstring(srcparts['word/styles.xml']);nst=E.fromstring(newparts['word/styles.xml'])
for st in nst:
 if st.get(qn('w:styleId'),'').startswith('Q4'):styles.append(deepcopy(st))
srcparts['word/document.xml']=E.tostring(sroot,xml_declaration=True,encoding='UTF-8',standalone=True)
srcparts['word/styles.xml']=E.tostring(styles,xml_declaration=True,encoding='UTF-8',standalone=True)
srcparts['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=E.fromstring(srcparts['[Content_Types].xml']);cnt='http://schemas.openxmlformats.org/package/2006/content-types'
if not any(x.get('Extension')=='png' for x in ct):E.SubElement(ct,'{'+cnt+'}Default',Extension='png',ContentType='image/png')
srcparts['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(FULL,'w',ZIP_DEFLATED) as z:
 for key,value in srcparts.items():z.writestr(key,value)
assert sha(SRC)==source_hash
assert all(E.tostring(a,method='c14n')==E.tostring(b,method='c14n') for a,b in zip(old[:866],list(body)[:866]))
# Editable text source and mapping retain all equation nodes and traceability.
latex=json.loads((WORK/'equations_latex.json').read_text())
md=[]
for c in content:
 if c['type']=='p' and not any(cc.get('caption')==c['text'] for cc in content if cc['type']=='table'):md.append(('## ' if c['style'] in ['Q4Title','Q4H2','Q4H3'] else '')+c['text'])
 elif c['type']=='formula':md.append('公式 '+','.join(c['numbers'])+'\n\n$$\n'+latex[str(c['source_body_index'])]+'\n$$')
 elif c['type']=='newformula':md.append('公式 '+','.join(c['numbers'])+'\n\n$$\n'+c['latex']+'\n$$')
 elif c['type']=='figure':md.append('!['+c['caption']+'](../../'+c['path']+')')
 elif c['type']=='table':
  md.append(c['caption']+'\n\n| '+' | '.join(c['headers']).replace('\n',' ')+' |\n| '+' | '.join(['---']*len(c['headers']))+' |\n'+'\n'.join('| '+' | '.join(map(str,row)).replace('\n',' ')+' |' for row in c['rows']))
(WORK/'重构正文.md').write_text('\n\n'.join(md))
(WORK/'content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2))
result={'source':str(SRC),'source_sha256':source_hash,'unchanged_pre_q4_body_elements':866,'preserved_display_equation_count':19,'added_algebraic_equations':2,'equation_map':eqmap,'model_recomputed':False,'outputs':{str(f.name):sha(f) for f in [OUT,FULL]},'source_model_code_reviewed':True,'full_model_rerun':False,'frozen_comparison':{'12m':main[0]-freeze[0],'24m':main[1]-freeze[1]},'input_files':{str(f.relative_to(ROOT)):sha(f) for f in (ROOT/'output_q4_evolution/tables').glob('*.csv')},'figures':{f:sha(ROOT/f) for f in used}}
(WORK/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps({'outputs':list(result['outputs']),'equations':sum(len(x) for x in eqmap.values()),'blocks':len(content)},ensure_ascii=False))
